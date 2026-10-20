"""
prf_helpers.py — ACME Medical PRF Processing Helpers
Document: WI-7.4.3-01 (dependency)

All utility functions for receiving inspection automation.
Import this module in any script Claude Code generates for this workflow.
"""

import io
import json
import os
import re
import shutil
import subprocess
import zipfile
from datetime import datetime
from pathlib import Path

import fitz  # PyMuPDF
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

# ─────────────────────────────────────────────
# TOOLS DIRECTORY — all temp/cache files live here, never in PO folders
# ─────────────────────────────────────────────
TOOLS_DIR = Path(__file__).parent

# Bump whenever this file changes with a substantive behavior fix — mirrors
# the date of this file's own most recent CLAUDE.md changelog entry. Added
# 2026-09-17 alongside write_manifest() so the manifest has something real
# to report for "prf_helpers.py version" (Step 9) instead of nothing.
PRF_HELPERS_VERSION = "2026-09-17"

# ─────────────────────────────────────────────
# TRACEABILITY MANIFEST (Step 9) — added 2026-09-17, resolves Open Item #8
# ─────────────────────────────────────────────
# WI-7.4.3-01 Step 9 has described a per-part JSON manifest as existing,
# mandatory output since at least Rev 01. It never existed — see CLAUDE.md
# Open Item #8 for the full history (confirmed missing 2026-07-22 and again
# 2026-09-11; TC-14 was incorrectly marked Pass against it on 7/28/2026,
# corrected 2026-09-17 alongside this fix).
#
# LOCATION: Step 9's literal text names `_Automation\logs\`, a shared Box
# folder. That folder tree was never built (Open Items #2/#3/#8: the
# `_Automation_Master` yearly-copy deployment was planned but never
# implemented) and this tool has only ever run from the engineer's own
# Desktop folder. LOGS_DIR therefore points at a `logs\` folder alongside
# this file — the same real-practice adaptation already made for FILE ACCESS
# RULES (2026-09-08 changelog) rather than a Box path that doesn't exist.
# This is a real, deliberate deviation from WI-7.4.3-01's literal wording,
# not an oversight, and needs Quality/Document Control to ratify it (or say
# where it should actually live) before Rev A — flagged again in the
# 2026-09-17 changelog entry.
LOGS_DIR = TOOLS_DIR / "logs"

_MANIFEST_NAME_RE = re.compile(r'[^A-Za-z0-9_.-]+')


def write_manifest(po_number: str, part_number: str, wi_revision: str, data: dict) -> Path:
    """
    Write one part's traceability manifest to LOGS_DIR, per WI-7.4.3-01 Step 9.

    Filename: <PO#>_<part#>_<YYYYMMDD-HHMM>_manifest.json, flat in LOGS_DIR
    (never inside a PO folder) — per Step 9, the filename alone must identify
    the run without opening the file, since logs are shared across every PO
    processed in a year.

    KNOWN GAP, not silently papered over: Step 9 also calls for recording
    "which method produced [each extracted fact] (pdfplumber vs. haiku vs.
    sonnet vs. cache hit)". That instrumentation does not exist anywhere in
    receive.py as of this writing — extraction call sites don't track which
    path they took. Rather than fabricate that field, `data` simply omits
    it; callers should not invent a value. See CLAUDE.md's 2026-09-17
    changelog entry — this is a real, separate, still-open piece of Step 9,
    not something this function silently resolved.

    Caller (receive.py's process_part()) supplies `data` with the real
    per-part content: Section 1 cross-check (Step 2.5) result, Section 2.1-
    2.6 results and the inputs behind them, Section 6 cross-reference
    findings, disposition, and every flag raised. This function does not
    validate field names — it trusts the caller for content, and only owns
    the file's location, naming, and the fields common to every manifest
    (timestamp, WI revision, this file's version).
    """
    LOGS_DIR.mkdir(exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d-%H%M")
    safe_po = _MANIFEST_NAME_RE.sub('_', str(po_number).strip()) or "UNKNOWNPO"
    safe_part = _MANIFEST_NAME_RE.sub('_', str(part_number).strip()) or "UNKNOWNPART"
    path = LOGS_DIR / f"{safe_po}_{safe_part}_{ts}_manifest.json"
    payload = {
        "timestamp": datetime.now().isoformat(timespec="seconds"),
        "wi_revision": wi_revision,
        "prf_helpers_version": PRF_HELPERS_VERSION,
        "extraction_method_tracking": "NOT INSTRUMENTED — see this function's docstring "
                                       "and CLAUDE.md Open Item #8 2026-09-17 entry",
        **data,
    }
    with open(path, "w", encoding="utf-8") as f:
        json.dump(payload, f, indent=2, default=str)
    return path


# ─────────────────────────────────────────────
# CHECKBOX HELPERS
# ─────────────────────────────────────────────

def _cb_font(cell):
    """Preserve the cell's existing font — never override it."""
    f = cell.font
    return Font(
        name=f.name if f.name else 'Calibri',
        size=f.size if f.size else 11,
        bold=f.bold if f.bold else False,
    )


def check_box(ws, coord, with_label=None, align=None):
    """Write ☑ to coord. with_label appends label text inline.

    Fixed 2026-09-17 (QE, direct report — recurring, permanent formatting
    bug, not a new template): the default alignment used to be 'center' only
    when there was NO label, and 'left' whenever with_label was given. Every
    real cell in the actual template bakes its label into the SAME cell as
    the checkbox (K30 "☐Yes", J36 "☐Raw Goods", I20 "☐Completed", L12/G3
    "☐N/A", etc. — see the 2026-07-27 Section 6 label-deletion changelog
    entry) and every one of those cells is itself center-aligned in the
    blank template. Left-aligning a labeled checkbox therefore visibly
    shoved the glyph+text over to the left edge of its field on every
    checked/unchecked write since this helper was written — exactly the
    "checkboxes and line items (6.2, N/A, Yes, No) drift to the side"
    symptom reported on real compiled packages. The handful of call sites
    that already passed align="center" explicitly (the Section 1/3/4
    rendering-only loops) were unaffected and are unchanged; this fixes the
    ones that relied on the default — I20, G3, L12, K30/N30, J36/N36.
    A caller can still force a different alignment by passing align=
    explicitly; only the default changed.
    """
    cell = ws[coord]
    cell.value = f'☑  {with_label}' if with_label else '☑'
    cell.font = _cb_font(cell)
    if align is None:
        align = 'center'
    cell.alignment = Alignment(horizontal=align, vertical='center')


def uncheck_box(ws, coord, with_label=None, align=None):
    """Write ☐ to coord — makes the empty box visible.

    See check_box()'s 2026-09-17 note — same default-alignment fix applies
    here (was 'left' whenever with_label was given; now always 'center'
    unless the caller passes align= explicitly).
    """
    cell = ws[coord]
    cell.value = f'☐  {with_label}' if with_label else '☐'
    cell.font = _cb_font(cell)
    if align is None:
        align = 'center'
    cell.alignment = Alignment(horizontal=align, vertical='center')


# ─────────────────────────────────────────────
# CHECKBOX INTEGRITY GUARD — added 2026-07-14
# ─────────────────────────────────────────────
# Root cause (see CLAUDE.md changelog, 2026-07-14): Section 1.2/1.3 checkboxes
# were going missing/wrong on repeated runs because the VML/ctrlProp state
# reader in receive.py silently returned {} (or a partial dict) on ANY parse
# failure, and the caller used vml_states.get(coord, False) — meaning "we
# couldn't read this cell" and "the engineer left it unchecked" were treated
# as identical. The automation would then happily write ☐ over a checkbox the
# engineer had actually checked, with only an easy-to-miss console warning.
#
# Fix: reads that can't be trusted now raise ChecklistCheckboxError instead of
# degrading to a default. This function is the second gate — it re-reads the
# cells right before save and confirms the glyph actually written matches what
# the automation intended, so a write-side or zip-surgery regression can't
# silently slip through either.


class ChecklistCheckboxError(Exception):
    """
    Raised when a PRF checkbox's state cannot be verified or does not match
    what the automation intended to write. Must never be caught and silently
    ignored — a PRF is a signed quality record, and defaulting an unverifiable
    checkbox to "unchecked" (or any other guessed value) is a data-integrity
    failure, not a recoverable one. Surface this to the operator every time.
    """
    pass


def verify_checkbox_fidelity(ws, expected: dict, sheet_label: str = ""):
    """
    Confirm every {coord: True/False} in `expected` actually landed in the
    worksheet as the correct glyph (☑ for True, ☐ for False) before the file
    is saved, exported, or compiled into a quality record.

    Call this immediately before save_preserving_drawings() with every
    checkbox coordinate the current run touched. Raises ChecklistCheckboxError
    on any mismatch instead of allowing a wrong/blank checkbox to reach
    In-Progress/ or Inspections/.
    """
    # 2026-07-27 fix (found while fixing Section 6.2's B31 label, PO26-1006):
    # this used to require the glyph at position 0 (val.startswith(...)).
    # Most checkbox cells are "☑ Label" (glyph first), but B31's pristine
    # format is "6.2\n☐ N/A" (a row/section number ahead of the glyph) —
    # a legitimate different cell layout, not a bug. A startswith check
    # rejected a correctly-written B31 every time, which would have blocked
    # every single run whenever that code path executed. Switched to a
    # containment check: still unambiguous, since a correctly-written cell
    # never contains both glyphs at once, and still catches a genuinely
    # wrong or missing glyph exactly as before.
    mismatches = []
    for coord, is_checked in expected.items():
        cell = ws[coord]
        val = str(cell.value or "")
        expected_glyph = '☑' if is_checked else '☐'
        wrong_glyph = '☐' if is_checked else '☑'
        if expected_glyph not in val or wrong_glyph in val:
            mismatches.append(f"{coord}: expected {expected_glyph!r}, found {val!r}")
    if mismatches:
        label = f" ({sheet_label})" if sheet_label else ""
        raise ChecklistCheckboxError(
            f"Checkbox fidelity check FAILED{label} — refusing to save/compile:\n  " +
            "\n  ".join(mismatches)
        )


# ─────────────────────────────────────────────
# EXCEL SAVE — PRESERVES DRAWINGS / VML
# ─────────────────────────────────────────────

def _zip_ext(name: str) -> str:
    """
    Lowercase extension of an OPC zip part name, '' if it has none.

    Deliberately not Path(name).suffix: OPC packages contain dot-leading part names
    such as "_rels/.rels", and pathlib treats ".rels" as a stem with NO suffix, which
    makes a content-type audit report perfectly valid parts as undeclared.
    """
    base = name.rsplit("/", 1)[-1]
    return base.rsplit(".", 1)[-1].lower() if "." in base else ""


def _reconcile_content_types(saved_ct: str, template_ct: str, final_names: set) -> str:
    """
    Merge into openpyxl's regenerated [Content_Types].xml the <Default>/<Override>
    declarations it dropped, but ONLY for parts that actually exist in final_names.

    Nothing gets declared that isn't in the package, and nothing in the package goes
    undeclared — an undeclared part or a dangling declaration is what makes Excel
    throw its "repair the workbook" prompt.
    """
    have_defaults = {e.lower() for e in
                     re.findall(r'<Default\s+Extension="([^"]+)"', saved_ct, re.I)}
    have_overrides = set(re.findall(r'<Override\s+PartName="([^"]+)"', saved_ct, re.I))
    tpl_defaults = dict(re.findall(
        r'<Default\s+Extension="([^"]+)"\s+ContentType="([^"]+)"\s*/>', template_ct, re.I))
    tpl_overrides = dict(re.findall(
        r'<Override\s+PartName="([^"]+)"\s+ContentType="([^"]+)"\s*/>', template_ct, re.I))

    used_exts = {_zip_ext(n) for n in final_names} - {""}
    add_defaults = {e: c for e, c in tpl_defaults.items()
                    if e.lower() in used_exts and e.lower() not in have_defaults}
    add_overrides = {p: c for p, c in tpl_overrides.items()
                     if p.lstrip("/") in final_names and p not in have_overrides}
    if not add_defaults and not add_overrides:
        return saved_ct

    inject = "".join(f'<Default Extension="{e}" ContentType="{c}"/>'
                     for e, c in sorted(add_defaults.items()))
    inject += "".join(f'<Override PartName="{p}" ContentType="{c}"/>'
                      for p, c in sorted(add_overrides.items()))
    m = re.search(r"<Types\b[^>]*>", saved_ct)
    if not m:
        return saved_ct
    return saved_ct[:m.end()] + inject + saved_ct[m.end():]


def save_preserving_drawings(wb, template_path: str, output_path: str):
    """
    Save workbook while restoring all drawing/VML/media files from the template.

    openpyxl strips VML drawing files, control properties, printer settings, embedded
    media, and worksheet relationship files during its normal save. This function
    re-injects them from the template so the xlsx structure remains intact.

    xl/media/ IS PRESERVED (fixed 2026-08-12, found on PO26-1011). It used not to be:
    the copy-forward filter matched "drawing"/"ctrlProp"/"printerSettings"/worksheet
    rels, which pulls in xl/drawings/_rels/vmlDrawing2.vml.rels — a relationship whose
    Target is ../media/image1.png, the ACME Medical letterhead logo — while "xl/media/"
    matched nothing. The saved workbook therefore shipped a DANGLING relationship, and
    because export_and_stamp_pdf() reads the logo out of the file handed to it as
    template_path, every PRF PDF exported from a just-saved workbook came out with the
    QF 7.4.3-02 header text and NO logo, silently. [Content_Types].xml is reconciled
    for the same reason: openpyxl regenerates it and drops the png/vml/ctrlProps
    declarations, leaving parts present but undeclared.

    NOTE on native form control checkboxes:
    Preserving the VML/ctrlProps PARTS is necessary but NOT sufficient (found
    2026-08-12, PO26-1012). openpyxl also rewrites every xl/worksheets/sheetN.xml
    and drops the elements that REFERENCE those parts — <drawing>, <legacyDrawing>,
    <legacyDrawingHF>, the <mc:AlternateContent><controls> block, and the r:id on
    <pageSetup>. The parts survive as orphans, so Excel shows no clickable
    checkboxes at all and garbage-collects the orphaned parts on its next save.
    ALWAYS follow this call with restore_form_controls(), which splices those
    references back in from the template. Then set answers with
    set_native_checkboxes() so the delivered xlsx stays editable.

    LibreOffice headless PDF export still cannot render VML form controls. Do not
    solve that by writing ☑/☐ glyphs into the delivered workbook — that is what
    made PRFs read-only. Render glyphs into a throwaway copy under TOOLS_DIR/_tmp
    and export the PDF from that copy instead.

    NEVER call wb.save(output_path) directly; always use this function.
    """
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    saved_bytes = buf.read()

    carry_files = {}
    with zipfile.ZipFile(template_path, "r") as tz:
        template_ct = tz.read("[Content_Types].xml").decode("utf-8")
        for name in tz.namelist():
            if ("drawing" in name.lower() or
                    "ctrlProp" in name or
                    "printerSettings" in name or
                    name.startswith("xl/media/") or
                    name.startswith("xl/worksheets/_rels/")):
                carry_files[name] = tz.read(name)

    final = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(saved_bytes), "r") as sz:
        existing_names = set(sz.namelist())
        saved_ct = sz.read("[Content_Types].xml").decode("utf-8")
        # Skip any file openpyxl already wrote — duplicate zip entries make Excel
        # prompt to repair the workbook.
        to_add = {n: d for n, d in carry_files.items() if n not in existing_names}
        new_ct = _reconcile_content_types(saved_ct, template_ct,
                                         existing_names | set(to_add))
        with zipfile.ZipFile(final, "w", zipfile.ZIP_DEFLATED) as oz:
            for item in sz.infolist():
                if item.filename == "[Content_Types].xml":
                    oz.writestr(item, new_ct.encode("utf-8"))
                else:
                    oz.writestr(item, sz.read(item.filename))
            for name, data in to_add.items():
                oz.writestr(name, data)

    with open(output_path, "wb") as f:
        f.write(final.getvalue())


# ─────────────────────────────────────────────
# NATIVE FORM CONTROL CHECKBOXES — added 2026-08-12 (PO26-1012)
# ─────────────────────────────────────────────
# Why this exists: the automation used to record every answer by writing a ☑/☐
# TEXT GLYPH into the checkbox's anchor cell. That renders correctly in the
# exported PDF, but it is not a checkbox — the engineer opening the delivered
# xlsx could no longer tick or untick anything, because openpyxl had stripped the
# real form controls out of the sheet XML (see save_preserving_drawings). A PRF
# that cannot be corrected in Excel is not an acceptable deliverable.
#
# The split of responsibility is now:
#   restore_form_controls()  — put the control wiring openpyxl removed back
#   set_native_checkboxes()  — set answers on the REAL controls (stays clickable)
#   glyph copy under _tmp    — text-glyph render used only to export the PDF

_SHEET_XML_RE = re.compile(r"xl/worksheets/sheet\d+\.xml$")

# Elements openpyxl drops from the worksheet, all of which sit between
# <pageSetup> and </worksheet> in schema order.
_TAIL_ELEMENT_RE = re.compile(
    r"<(?:drawing|legacyDrawing|legacyDrawingHF|picture|oleObjects|controls)\b"
    r"|<mc:AlternateContent\b")
_DROP_FROM_SAVED_RE = re.compile(
    r"<drawing\b[^>]*/>"
    r"|<legacyDrawing(?:HF)?\b[^>]*/>"
    r"|<mc:AlternateContent\b.*?</mc:AlternateContent>"
    r"|<oleObjects\b.*?</oleObjects>",
    re.S)


def _sheet_xml_names(names):
    return sorted(n for n in names if _SHEET_XML_RE.match(n))


def _anchor_coord(ws, col_idx, row_idx):
    """0-indexed VML (col, row) -> writable coord, resolved to its merge anchor."""
    from openpyxl.utils import get_column_letter, range_boundaries
    coord = f"{get_column_letter(col_idx + 1)}{row_idx + 1}"
    for rng in ws.merged_cells.ranges:
        if coord in rng:
            min_col, min_row, _, _ = range_boundaries(str(rng))
            return f"{get_column_letter(min_col)}{min_row}"
    return coord


def restore_form_controls(xlsx_path: str, template_path: str):
    """
    Re-attach the native form controls openpyxl orphaned, so the saved PRF opens
    in Excel with working, clickable checkboxes.

    save_preserving_drawings() carries the vmlDrawing/ctrlProps PARTS forward, but
    openpyxl's regenerated sheet XML no longer points at them. This splices the
    template's trailing <drawing>/<legacyDrawing>/<legacyDrawingHF>/<controls>
    elements and the <pageSetup r:id> back into each saved sheet, re-declares any
    XML namespace those elements need (xdr, x14, mc), and forces the template's
    worksheet .rels so every restored r:id resolves.

    Call immediately after save_preserving_drawings(). Raises
    ChecklistCheckboxError if the saved workbook holds a worksheet relationship
    the template does not, since template-wins would silently discard it.
    """
    with zipfile.ZipFile(template_path) as tz:
        tpl = {n: tz.read(n) for n in tz.namelist()}
    with zipfile.ZipFile(xlsx_path) as sz:
        saved = {n: sz.read(n) for n in sz.namelist()}
        order = list(sz.namelist())

    restored = 0
    for sx in _sheet_xml_names(tpl):
        if sx not in saved:
            continue
        tpl_body = tpl[sx].decode("utf-8")
        saved_body = saved[sx].decode("utf-8")

        # Locate the template's trailing block (everything after <pageSetup>).
        anchor = tpl_body.rfind("<pageSetup")
        if anchor < 0:
            anchor = max(tpl_body.rfind("</mergeCells>"), 0)
        m = _TAIL_ELEMENT_RE.search(tpl_body, anchor)
        if not m:
            continue                     # sheet genuinely has no controls
        end = tpl_body.rfind("</worksheet>")
        tail = tpl_body[m.start():end]

        # Drop anything openpyxl wrote in that space so nothing is duplicated.
        saved_body = _DROP_FROM_SAVED_RE.sub("", saved_body)

        # <pageSetup> loses its r:id (the printerSettings link) on the way through
        # openpyxl; take the template's element verbatim.
        tpl_ps = re.search(r"<pageSetup\b[^>]*/?>", tpl_body)
        if tpl_ps:
            saved_body = re.sub(r"<pageSetup\b[^>]*/?>", lambda _: tpl_ps.group(0),
                                saved_body, count=1)

        # The restored markup uses xdr:/x14:/mc: prefixes — declare any the saved
        # root is missing, or Excel rejects the file as malformed.
        tpl_root = re.search(r"<worksheet\b[^>]*>", tpl_body).group(0)
        saved_root = re.search(r"<worksheet\b[^>]*>", saved_body).group(0)
        add = [f'{k}="{v}"' for k, v in re.findall(r'(xmlns:[\w]+)="([^"]+)"', tpl_root)
               if k not in saved_root]
        ignorable = re.search(r'mc:Ignorable="([^"]+)"', tpl_root)
        if ignorable and "mc:Ignorable" not in saved_root:
            add.append(ignorable.group(0))
        if add:
            new_root = saved_root[:-1].rstrip() + " " + " ".join(add) + ">"
            saved_body = saved_body.replace(saved_root, new_root, 1)

        cut = saved_body.rfind("</worksheet>")
        saved[sx] = (saved_body[:cut] + tail + saved_body[cut:]).encode("utf-8")
        restored += 1

        # Force the template's worksheet rels so every restored r:id resolves.
        rels = sx.replace("worksheets/", "worksheets/_rels/") + ".rels"
        if rels in tpl:
            tpl_targets = set(re.findall(r'Target="([^"]+)"',
                                         tpl[rels].decode("utf-8")))
            if rels in saved:
                extra = [t for t in re.findall(r'Target="([^"]+)"',
                                               saved[rels].decode("utf-8"))
                         if t not in tpl_targets]
                if extra:
                    raise ChecklistCheckboxError(
                        f"{Path(xlsx_path).name}: {rels} holds relationship(s) the "
                        f"template does not — restoring the template's rels would "
                        f"discard them: {extra}")
            saved[rels] = tpl[rels]
            if rels not in order:
                order.append(rels)

    with zipfile.ZipFile(xlsx_path, "w", zipfile.ZIP_DEFLATED) as oz:
        for name in order:
            oz.writestr(name, saved[name])
    return restored


def _control_map(xlsx_path, wb):
    """
    {sheet_name: {coord: {"vml_id", "ctrl_prop", "checked"}}} for every native
    checkbox.

    The two halves are joined on the ANCHOR CELL, not on the shape id. Both the
    sheet XML <control> element and the VML <v:shape> carry a from-col/from-row
    anchor, and those always agree — whereas <control shapeId="5"> and VML
    id="_x0000_s1029" do NOT have to match. Joining on id silently dropped 4 of
    the 21 Section 1-2 checkboxes (D8/G8/J8/L24, the oldest controls in the
    template), which is exactly the kind of silent partial read that
    ChecklistCheckboxError exists to prevent.
    """
    out = {}
    with zipfile.ZipFile(xlsx_path) as z:
        names = z.namelist()
        for idx, sx in enumerate(_sheet_xml_names(names)):
            if idx >= len(wb.sheetnames):
                break
            sheet_name = wb.sheetnames[idx]
            ws = wb[sheet_name]
            body = z.read(sx).decode("utf-8", "replace")
            rels_name = sx.replace("worksheets/", "worksheets/_rels/") + ".rels"
            if rels_name not in names:
                continue
            rels = z.read(rels_name).decode("utf-8", "replace")
            rmap = dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', rels))
            # Take the VML that <legacyDrawing> points at, NOT just any vmlDrawing
            # target: a sheet with a header/footer image also carries a
            # <legacyDrawingHF> VML, which holds no checkbox shapes at all.
            leg = re.search(r'<legacyDrawing\s+r:id="([^"]+)"', body)
            vml_target = rmap.get(leg.group(1)) if leg else None
            if not vml_target:
                vml_target = next((t for t in rmap.values() if "vmlDrawing" in t), None)
            if not vml_target or "vmlDrawing" not in vml_target:
                continue
            vml_path = "xl/drawings/" + vml_target.split("/")[-1]
            if vml_path not in names:
                continue
            vml = z.read(vml_path).decode("utf-8", "replace")

            # --- VML half: anchor cell -> (shape id attribute, checked state) ---
            by_coord = {}
            for shape in re.findall(r"<v:shape\b.*?</v:shape>", vml, re.S):
                if 'ObjectType="Checkbox"' not in shape:
                    continue
                sid = re.search(r'<v:shape\b[^>]*\bid="([^"]+)"', shape)
                anc = re.search(r"<x:Anchor>\s*([^<]+)</x:Anchor>", shape)
                if not sid or not anc:
                    continue
                nums = [int(x.strip()) for x in anc.group(1).split(",")]
                by_coord[_anchor_coord(ws, nums[0], nums[2])] = (
                    sid.group(1),
                    bool(re.search(r"<x:Checked>\s*1\s*</x:Checked>", shape)))

            # --- sheet XML half: anchor cell -> ctrlProps part ---
            per_sheet = {}
            for ctrl in re.findall(r"<control\b.*?</control>|<control\b[^>]*/>",
                                   body, re.S):
                rid = re.search(r'r:id="(rId\d+)"', ctrl)
                frm = re.search(r"<from>\s*<xdr:col>(\d+)</xdr:col>.*?"
                                r"<xdr:row>(\d+)</xdr:row>", ctrl, re.S)
                if not rid or not frm:
                    continue
                coord = _anchor_coord(ws, int(frm.group(1)), int(frm.group(2)))
                if coord not in by_coord or coord in per_sheet:
                    continue          # no VML twin, or an mc:Fallback duplicate
                vml_id, checked = by_coord[coord]
                per_sheet[coord] = {
                    "vml_id": vml_id,
                    "ctrl_prop": "xl/" + rmap.get(rid.group(1), "").replace("../", ""),
                    "vml": vml_path,
                    "checked": checked,
                }
            if len(per_sheet) != len(by_coord):
                missed = sorted(set(by_coord) - set(per_sheet))
                raise ChecklistCheckboxError(
                    f"{Path(xlsx_path).name} [{sheet_name}]: {len(missed)} VML "
                    f"checkbox(es) have no matching <control> element and would be "
                    f"invisible to the automation: {missed}")
            out[sheet_name] = per_sheet
    return out


def read_native_checkboxes(xlsx_path: str, wb=None):
    """{sheet_name: {anchor_coord: bool}} for every live native checkbox."""
    from openpyxl import load_workbook
    own = wb is None
    if own:
        wb = load_workbook(xlsx_path, rich_text=True)
    try:
        return {s: {c: v["checked"] for c, v in m.items()}
                for s, m in _control_map(xlsx_path, wb).items()}
    finally:
        if own:
            wb.close()


def set_native_checkboxes(xlsx_path: str, states: dict):
    """
    Set real form-control checkbox states in place: states = {sheet_name:
    {anchor_coord: True/False}}. The boxes stay fully clickable in Excel.

    Writes BOTH representations Office keeps in sync — <x:Checked> in the VML
    (what LibreOffice and the legacy renderer read) and checked="Checked" in the
    ctrlProps part (what Excel reads). Raises ChecklistCheckboxError if a
    requested coordinate has no native checkbox, rather than silently no-op'ing
    an answer on a signed quality record.
    """
    from openpyxl import load_workbook
    wb = load_workbook(xlsx_path, rich_text=True)
    try:
        cmap = _control_map(xlsx_path, wb)
    finally:
        wb.close()

    missing = [f"{s}!{c}" for s, want in states.items()
               for c in want if c not in cmap.get(s, {})]
    if missing:
        raise ChecklistCheckboxError(
            f"{Path(xlsx_path).name}: no native checkbox found at {missing} — "
            f"refusing to record an answer that would not appear in the workbook.")

    with zipfile.ZipFile(xlsx_path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
        order = list(z.namelist())

    changed = 0
    for sheet, want in states.items():
        for coord, checked in want.items():
            info = cmap[sheet][coord]
            # --- VML: <x:Checked>1</x:Checked> inside the shape's ClientData ---
            vml = parts[info["vml"]].decode("utf-8")
            pat = re.compile(
                r'(<v:shape\b[^>]*\bid="' + re.escape(info["vml_id"]) +
                r'"[^>]*>.*?</v:shape>)', re.S)
            m = pat.search(vml)
            if not m:
                raise ChecklistCheckboxError(
                    f"{Path(xlsx_path).name}: VML shape for {sheet}!{coord} vanished.")
            shape = re.sub(r"\s*<x:Checked>\s*\d+\s*</x:Checked>", "", m.group(1))
            if checked:
                if "<x:NoThreeD" in shape:
                    shape = shape.replace("<x:NoThreeD",
                                          "<x:Checked>1</x:Checked>\n   <x:NoThreeD", 1)
                else:
                    shape = shape.replace("</x:ClientData>",
                                          "<x:Checked>1</x:Checked>\n  </x:ClientData>", 1)
            parts[info["vml"]] = (vml[:m.start(1)] + shape + vml[m.end(1):]).encode("utf-8")

            # --- ctrlProps: checked="Checked" attribute on <formControlPr> ---
            cp_name = info["ctrl_prop"]
            if cp_name in parts:
                cp = parts[cp_name].decode("utf-8")
                cp = re.sub(r'\s+checked="[^"]*"', "", cp)
                if checked:
                    cp = cp.replace("<formControlPr ",
                                    '<formControlPr checked="Checked" ', 1)
                parts[cp_name] = cp.encode("utf-8")
            changed += 1

    with zipfile.ZipFile(xlsx_path, "w", zipfile.ZIP_DEFLATED) as oz:
        for name in order:
            oz.writestr(name, parts[name])
    return changed


def verify_native_checkboxes(xlsx_path: str, expected: dict):
    """
    Post-save gate: confirm every {sheet: {coord: bool}} actually reads back off
    the real controls in the saved file. Same contract as
    verify_checkbox_fidelity(), but for native checkboxes instead of glyphs.
    """
    actual = read_native_checkboxes(xlsx_path)
    problems = []
    for sheet, want in expected.items():
        got = actual.get(sheet, {})
        for coord, state in want.items():
            if coord not in got:
                problems.append(f"{sheet}!{coord}: no native checkbox present")
            elif got[coord] != state:
                problems.append(f"{sheet}!{coord}: expected "
                                f"{'checked' if state else 'unchecked'}, "
                                f"file has {'checked' if got[coord] else 'unchecked'}")
    if problems:
        raise ChecklistCheckboxError(
            f"{Path(xlsx_path).name} native checkbox verification FAILED:\n  "
            + "\n  ".join(problems))


def strip_orphan_checkbox(xlsx_path: str, sheet_name: str, coord: str) -> bool:
    """
    Permanently remove one native checkbox control — both halves, the VML
    <v:shape> and its <control> twin in the sheet XML — from a saved PRF.

    Added 2026-09-17 to actually close Open Item #11 (G2 on the Sections
    4-6 sheet). The 2026-07-27 fix only ever cleared a stray TEXT value a
    prior run had written into G2's cell (`_S46_IGNORED_CHECKBOXES` in
    receive.py). It never touched the native checkbox shape itself, which
    is a separate object anchored at that cell and renders in LibreOffice's
    PDF export on its own, regardless of what text (if any) the cell holds.
    Confirmed directly on a real, already-issued compiled PRF (SR90-1005,
    PO26-1010, well after the July fix): a small, empty, unlabeled checkbox
    image sits next to Section 4's N/A label on every PRF this tool has
    ever compiled, because nothing had ever removed the underlying shape.

    Removes BOTH halves together, deliberately. `_control_map()` — which
    `read_native_checkboxes()`/`set_native_checkboxes()`/
    `verify_native_checkboxes()` all depend on — raises
    ChecklistCheckboxError if a VML shape has no matching <control> twin,
    on the reasoning that a checkbox visible to Excel but invisible to this
    tool is worse than either extreme (see that function's docstring).
    Removing only one half here would trip that exact guard on every
    subsequent run against this file. Leaves the ctrlProps part itself
    orphaned in the package — unreferenced, and harmless: unlike an
    unpaired VML shape, nothing in this codebase's integrity checks looks
    for an orphaned ctrlProps part, and Excel/LibreOffice both silently
    ignore a part nothing points to.

    Call AFTER restore_form_controls() (which would otherwise splice the
    control straight back in from the template) and BEFORE
    set_native_checkboxes()/verify_native_checkboxes() for that sheet.
    Returns False and does nothing if the coordinate has no live control
    left (e.g. called twice) — not an error, since the end state is the
    same either way. Raises ChecklistCheckboxError rather than silently
    leaving a half-removed checkbox if the VML half and sheet-XML half
    don't agree on what's there — the same fail-loud posture as every
    other native-checkbox function in this file. Does not touch the saved
    file at all unless both halves are found and removed successfully.
    """
    from openpyxl import load_workbook
    wb = load_workbook(xlsx_path, rich_text=True)
    try:
        cmap = _control_map(xlsx_path, wb)
        ws = wb[sheet_name]
        sheet_idx = wb.sheetnames.index(sheet_name)
    finally:
        wb.close()

    info = cmap.get(sheet_name, {}).get(coord)
    if info is None:
        return False  # already gone — nothing to do, not an error

    with zipfile.ZipFile(xlsx_path) as z:
        names = z.namelist()
        parts = {n: z.read(n) for n in names}
        order = list(names)

    # --- VML half: drop the whole <v:shape id="...">...</v:shape> block ---
    vml_name = info["vml"]
    vml = parts[vml_name].decode("utf-8")
    vml_pat = re.compile(
        r'\s*<v:shape\b[^>]*\bid="' + re.escape(info["vml_id"]) + r'"[^>]*>.*?</v:shape>',
        re.S)
    m = vml_pat.search(vml)
    if not m:
        raise ChecklistCheckboxError(
            f"{Path(xlsx_path).name}: expected VML shape {info['vml_id']} for "
            f"{sheet_name}!{coord} not found — refusing to guess; file left unchanged.")
    new_vml = vml[:m.start()] + vml[m.end():]

    # --- sheet XML half: drop the <control> element anchored at this coord ---
    sx = _sheet_xml_names(names)[sheet_idx]
    body = parts[sx].decode("utf-8")
    removed = None
    for ctrl in re.findall(r"<control\b.*?</control>|<control\b[^>]*/>", body, re.S):
        rid = re.search(r'r:id="(rId\d+)"', ctrl)
        frm = re.search(r"<from>\s*<xdr:col>(\d+)</xdr:col>.*?"
                        r"<xdr:row>(\d+)</xdr:row>", ctrl, re.S)
        if not rid or not frm:
            continue
        if _anchor_coord(ws, int(frm.group(1)), int(frm.group(2))) == coord:
            removed = ctrl
            break
    if removed is None:
        raise ChecklistCheckboxError(
            f"{Path(xlsx_path).name}: expected <control> element for "
            f"{sheet_name}!{coord} not found in sheet XML — refusing to leave "
            f"a half-removed checkbox; file left unchanged.")
    new_body = body.replace(removed, "", 1)

    # Only mutate the in-memory parts dict once both halves are confirmed —
    # the on-disk file is untouched until the write below, so either
    # exception above leaves the saved PRF exactly as it was.
    parts[vml_name] = new_vml.encode("utf-8")
    parts[sx] = new_body.encode("utf-8")

    with zipfile.ZipFile(xlsx_path, "w", zipfile.ZIP_DEFLATED) as oz:
        for name in order:
            oz.writestr(name, parts[name])
    return True


def verify_xlsx_package_integrity(xlsx_path: str):
    """
    Confirm a saved PRF is a structurally valid OPC package: every relationship
    Target resolves to a part that exists, and every part is declared in
    [Content_Types].xml.

    A dangling relationship or an undeclared part is what makes Excel throw its
    "we found a problem with some content — do you want us to repair" prompt on a
    signed quality record. Call this after save_preserving_drawings(). Raises
    ChecklistCheckboxError (the module's data-integrity error) listing every problem.
    """
    with zipfile.ZipFile(xlsx_path) as z:
        names = set(z.namelist())
        ct = z.read("[Content_Types].xml").decode("utf-8")
        dangling = []
        for rels_name in [n for n in names if n.endswith(".rels")]:
            # A .rels file at a/b/_rels/c.xml.rels describes targets relative to a/b.
            base = "/".join(rels_name.split("/")[:-2])
            body = z.read(rels_name).decode("utf-8", "replace")
            for target, mode in re.findall(
                    r'Target="([^"]+)"(?:\s+TargetMode="([^"]+)")?', body):
                if mode == "External" or target.startswith(("http", "mailto", "/")):
                    continue
                parts = [p for p in base.split("/") if p]
                for seg in target.split("/"):
                    if seg == "..":
                        if parts:
                            parts.pop()
                    elif seg not in ("", "."):
                        parts.append(seg)
                resolved = "/".join(parts)
                if resolved not in names:
                    dangling.append(f"{rels_name} -> {target}")

    defaults = {e.lower() for e in
                re.findall(r'<Default\s+Extension="([^"]+)"', ct, re.I)}
    overrides = {p.lstrip("/") for p in
                 re.findall(r'<Override\s+PartName="([^"]+)"', ct, re.I)}
    undeclared = [n for n in sorted(names)
                  if n != "[Content_Types].xml" and n not in overrides
                  and _zip_ext(n) not in defaults]

    problems = ([f"dangling relationship: {d}" for d in dangling] +
                [f"part not declared in [Content_Types].xml: {u}" for u in undeclared])
    if problems:
        raise ChecklistCheckboxError(
            f"{Path(xlsx_path).name} OPC package integrity check FAILED — refusing to "
            f"export/compile:\n  " + "\n  ".join(problems))


# ─────────────────────────────────────────────
# PDF EXPORT + HEADER STAMP
# ─────────────────────────────────────────────

def _load_font(windows_filename, fitz_fallback):
    """Load a font by Windows filename, fall back to a built-in fitz font."""
    win_path = os.path.join(r'C:\Windows\Fonts', windows_filename)
    if os.path.exists(win_path):
        return fitz.Font(fontfile=win_path)
    return fitz.Font(fontname=fitz_fallback)


# Fixed, known-good fallback source for the ACME Medical letterhead logo —
# added 2026-09-11 (WI-7.4.3-01 gap analysis) to harden export_and_stamp_pdf()
# against a stripped xl/media/ in the PRF being exported (per its own comment
# below, save_preserving_drawings()'s zip surgery — or, before the 2026-09-11
# restore_form_controls() fix, an Excel open/save that pruned the part
# outright — can leave a PRF with no logo image at all).
#
# CONFIRMED 2026-09-11 directly against the real Desktop tools folder (via
# the device bridge, `device_list_dir` on `templates\`): the ONLY file
# present there is "PRF template CONTROLS.xlsx" — not a file named after
# QF 7.4.3-02 as this constant originally guessed. Corrected to match the
# real filename. Confirm this file's own xl/media/ actually contains the
# ACME Medical logo before relying on it in production — this session could
# confirm the filename but not open the file itself.
FALLBACK_LOGO_TEMPLATE = TOOLS_DIR / "templates" / "PRF template CONTROLS.xlsx"


class LogoMissingError(Exception):
    """
    Raised when export_and_stamp_pdf() cannot find a ACME Medical letterhead
    logo image in EITHER the PRF being exported OR the fixed fallback
    template — refusing to silently ship a headerless PDF onto a signed
    quality record. See FALLBACK_LOGO_TEMPLATE above.
    """
    pass


class TextClippingError(Exception):
    """
    Raised (or, for the soft-check caller in receive.py, converted to a
    flag) when a merged-cell text field's rendered content was cut off or
    visually bisected by its own row's bottom rule in an exported PDF page.
    See check_row_text_clipping() docstring for the full context.
    """
    pass


def check_row_text_clipping(pdf_path: str, page_index: int, cell_rect,
                            field_label: str = "") -> list:
    """
    Detect a text field whose rendered content overflowed its own row and
    was clipped, or worse, visually bisected by the row's bottom rule, in an
    exported PDF page. Added 2026-09-11 (Priority 3 item #6, WI-7.4.3-01 gap
    analysis 2026-09-11).

    Root cause this addresses: LibreOffice's headless PDF export can silently
    drop overflow text past a fixed-height row, or draw straight through it
    at the row's bottom rule — while the source XLSX's own text layer still
    holds the complete, unclipped note. A text-only "did the full note make
    it into the exported file" check can therefore PASS on a rendered page a
    human would see as visibly cut in half. Found on real Section 4
    (DVT/acceptance-comment) notes on the SterilizerCo POs (PO26-1005/
    PO26-1009) and worked around ad hoc in bespoke per-PO scripts; this is
    the general-purpose port of that same geometric idea into the shared
    helper module, so any caller — not only a bespoke script — can apply it
    to any merged text field.

    Geometric approach: `cell_rect` is the field's own rectangle **on the
    rendered PDF page**, in PDF points (fitz.Rect or a 4-tuple) — e.g. the
    row's own vector-drawn border, read from `page.get_drawings()` the same
    way `_find_stadium_oval_coords()` in receive.py locates a hand-drawn
    oval, NOT a value back-computed from the XLSX's row-height metadata
    (LibreOffice's "fit to page" scaling can silently change the effective
    row height at export time, which would make an Excel-side computation
    unreliable). This function renders `page_index`'s text spans, keeps only
    spans whose bounding box overlaps `cell_rect` horizontally, and unions
    their vertical extent into one text-block rectangle. It flags when that
    text block's own bottom edge falls AT OR BELOW `cell_rect`'s bottom edge
    (with a small tolerance for normal cell padding) — i.e. the text needed
    more vertical room than the row actually has, the same condition a human
    reviewer would see as clipped or struck through by the row's own rule.

    Returns a list with zero or one flag string (never raises — a
    supplementary check, same "does not block the run on failure" posture
    as check_document_completeness()/check_calibration_date() in
    receive.py). An unreadable page or an empty/degenerate cell_rect
    degrades to no finding rather than a false positive.

    **NOT YET VERIFIED against a real filled-out Section 4 note or a real
    exported QF 7.4.3-02 page** — this is a from-scratch reconstruction of
    the geometric idea described in the gap analysis, not a port of the
    bespoke per-PO script's own already-proven implementation (that script
    was not available to this session to copy from directly). Treat any
    flag or silence from this function as a hypothesis pending confirmation
    against a real exported PDF, the same caveat this codebase already
    attaches to a first-pass geometric guess elsewhere (e.g. Open Item #17's
    original calibration-cell reference) — see the corresponding CLAUDE.md
    Open Item. No caller in this codebase invokes this function yet with a
    real `cell_rect` for Section 4 specifically; wiring it in requires first
    confirming, against a real exported PDF, where Section 4's comment field
    actually renders on the page.
    """
    try:
        r = fitz.Rect(cell_rect)
        if r.is_empty or r.is_infinite:
            return []
        doc = fitz.open(pdf_path)
        if not (0 <= page_index < doc.page_count):
            doc.close()
            return []
        page = doc[page_index]
        spans_bottom = None
        for block in page.get_text("dict").get("blocks", []):
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    sx0, sy0, sx1, sy1 = span["bbox"]
                    # Keep only spans whose horizontal extent actually overlaps
                    # the field's own column range — text from a neighboring
                    # column sitting on a similar vertical band must not count.
                    if sx1 < r.x0 - 1 or sx0 > r.x1 + 1:
                        continue
                    if sy0 < r.y0 - 2:
                        # Above the field entirely (e.g. the row's own label) —
                        # not part of this field's rendered content.
                        continue
                    spans_bottom = sy1 if spans_bottom is None else max(spans_bottom, sy1)
        doc.close()
        if spans_bottom is None:
            return []
        _tolerance = 1.5  # points — ordinary bottom cell padding, not a defect
        if spans_bottom >= r.y1 - _tolerance:
            label = f" ({field_label})" if field_label else ""
            return [
                f"TEXT CLIPPING RISK{label}: rendered text bottom "
                f"({spans_bottom:.1f}pt) reaches or passes the field's own "
                f"row boundary ({r.y1:.1f}pt) on page {page_index + 1} of "
                f"{Path(pdf_path).name} — this note may be cut off or drawn "
                f"through by the row's bottom rule. Open the exported PDF and "
                f"confirm the full note is visible; if not, the row needs more "
                f"height or the note needs to be shortened before signing."
            ]
        return []
    except Exception as e:
        # Never let a supplementary geometric check crash export — degrade to
        # "nothing to report," same posture as every other soft check here.
        print(f"  [WARN] check_row_text_clipping: could not evaluate ({e}) — skipped.")
        return []


def export_and_stamp_pdf(xlsx_path: str, out_dir: str, template_path: str) -> str:
    """
    Export xlsx to PDF via LibreOffice headless, then stamp logo + header on all pages.
    Returns the path to the stamped PDF, or None on failure.

    Logo sourcing (hardened 2026-09-11 — WI-7.4.3-01 gap analysis, Priority
    3 item 18): previously, if template_path's own xl/media/ had no image —
    which can happen after save_preserving_drawings()'s zip surgery, or (pre-
    2026-09-11) after Excel silently pruned an orphaned media part on open+
    save — this printed a WARNING and shipped a headerless PDF onto a signed
    quality record, with nothing stopping the run. Now falls back to
    FALLBACK_LOGO_TEMPLATE's own xl/media/ if template_path has none, and
    raises LogoMissingError (loud, blocks the run) if NEITHER source has a
    usable image, instead of silently degrading.
    """
    lo_candidates = [
        "libreoffice",
        "soffice",
        str(TOOLS_DIR / "LibreOffice" / "program" / "soffice.exe"),  # tools-folder install
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]
    for lo in lo_candidates:
        if shutil.which(lo) or Path(lo).exists():
            result = subprocess.run(
                [lo, "--headless", "--convert-to", "pdf", "--outdir", out_dir, xlsx_path],
                capture_output=True,
                timeout=60,
            )
            if result.returncode != 0:
                print("LibreOffice export failed:", result.stderr.decode())
                return None
            break
    else:
        print("LibreOffice not found. Install from https://libreoffice.org")
        return None

    raw_pdf = str(Path(out_dir) / (Path(xlsx_path).stem + ".pdf"))

    # Extract logo — scan xl/media/ for any image rather than assuming image1.png
    def _find_logo(xlsx_source) -> bytes | None:
        try:
            with zipfile.ZipFile(xlsx_source, 'r') as z:
                media = [n for n in z.namelist()
                         if n.startswith('xl/media/')
                         and n.lower().endswith(('.png', '.jpg', '.jpeg', '.gif', '.bmp', '.emf'))]
                if media:
                    return z.read(media[0])
        except (OSError, zipfile.BadZipFile):
            pass
        return None

    logo_bytes = _find_logo(template_path)
    logo_source = template_path
    if logo_bytes is None and FALLBACK_LOGO_TEMPLATE.exists():
        logo_bytes = _find_logo(FALLBACK_LOGO_TEMPLATE)
        logo_source = FALLBACK_LOGO_TEMPLATE
    if logo_bytes is None:
        # 2026-09-11: previously printed a WARNING and shipped a letterhead-
        # less quality record. That silent degradation is exactly what this
        # hardening pass exists to close — fail loudly instead, per Task #18.
        raise LogoMissingError(
            f"No ACME Medical letterhead logo image found in {Path(template_path).name}'s "
            f"xl/media/, and the fallback template ({FALLBACK_LOGO_TEMPLATE}) is either "
            f"missing or also has no usable image. Refusing to export a headerless PDF "
            f"onto a signed quality record. Confirm FALLBACK_LOGO_TEMPLATE points at a "
            f"real, pristine QF 7.4.3-02 master copy with an intact xl/media/, or repair "
            f"{Path(template_path).name}'s own media before re-running."
        )
    if logo_source != template_path:
        print(f"  [WARN] {Path(template_path).name} had no logo in xl/media/ — used the "
              f"fallback template ({FALLBACK_LOGO_TEMPLATE.name}) instead. This PRF's own "
              f"media was likely stripped by an earlier save — investigate before this "
              f"recurs (see CLAUDE.md Open Item, restore_form_controls()).")

    font_title  = _load_font('comicbd.ttf', 'helv')
    font_docnum = _load_font('times.ttf', 'Times-Roman')

    TITLE = "Quality Product Release Form"
    DOCNUM_LINES = [
        "Document Number: QF 7.4.3-02",
        "Date: 01/2026",
        "Revision E",
    ]

    tmp = raw_pdf + ".tmp"
    shutil.copy(raw_pdf, tmp)
    doc = fitz.open(tmp)

    for page in doc:
        pw = page.rect.width
        page.add_redact_annot(fitz.Rect(130, 0, pw, 58), fill=(1, 1, 1))
        page.apply_redactions()
        if logo_bytes:
            page.insert_image(fitz.Rect(28, 8, 133, 46), stream=logo_bytes)

        writer = fitz.TextWriter(page.rect, color=(0, 0, 0))
        tw = font_title.text_length(TITLE, fontsize=13.2)
        writer.append(((pw / 2) - (tw / 2), 30), TITLE, font=font_title, fontsize=13.2)

        x_right = pw - 5
        for i, line in enumerate(DOCNUM_LINES):
            lw = font_docnum.text_length(line, fontsize=7.92)
            writer.append((x_right - lw, 18 + i * 8.5), line,
                          font=font_docnum, fontsize=7.92)
        writer.write_text(page)

    doc.save(raw_pdf, garbage=4, deflate=True)
    doc.close()
    os.unlink(tmp)
    return raw_pdf


# ─────────────────────────────────────────────
# PDF COMPILATION + GHOSTSCRIPT RE-DISTILL
# ─────────────────────────────────────────────

def _find_ghostscript() -> str:
    """Locate the Ghostscript executable across platforms."""
    candidates = [
        "gswin64c",
        "gswin32c",
        "gs",
        r"C:\Program Files\gs\gs10.04.0\bin\gswin64c.exe",
        r"C:\Program Files\gs\gs10.03.1\bin\gswin64c.exe",
        r"C:\Program Files\gs\gs10.02.1\bin\gswin64c.exe",
        r"C:\Program Files (x86)\gs\gs10.04.0\bin\gswin32c.exe",
    ]
    # Auto-detect any gs* subfolder in the tools directory (e.g. gs10.07.1\bin\gswin64c.exe)
    for gs_dir in sorted(TOOLS_DIR.glob("gs*"), reverse=True):  # newest version first
        for exe_name in ["gswin64c.exe", "gswin32c.exe"]:
            candidate = gs_dir / "bin" / exe_name
            if candidate.exists():
                candidates.insert(0, str(candidate))
    for c in candidates:
        if shutil.which(c) or Path(c).exists():
            return c
    return None


def _gs_redistill(input_path: str, output_path: str) -> bool:
    """
    Re-distill a PDF through Ghostscript to produce a clean, fully editable PDF.
    Strips all security dictionaries, digital signature flags, encryption, and
    AcroForm locks. Returns True on success, False if Ghostscript is unavailable.
    """
    gs = _find_ghostscript()
    if gs is None:
        return False
    cmd = [
        gs, "-dBATCH", "-dNOPAUSE", "-dQUIET", "-dSAFER",
        "-sDEVICE=pdfwrite", "-dCompatibilityLevel=1.7",
        "-dPDFSETTINGS=/prepress", "-dNOENCRYPT",
        f"-sOutputFile={output_path}", input_path,
    ]
    result = subprocess.run(cmd, capture_output=True, timeout=120)
    return result.returncode == 0


class PdfEditabilityError(Exception):
    """
    Raised when a compiled receiving package cannot be signed or page-edited.

    Every compiled package must be signable and page-editable: the Quality Manager
    signs it, and pages get replaced when a supplier document is corrected. A locked
    package is a failed deliverable, so this must never be caught and ignored.
    """
    pass


# Catalog keys that lock editing/signing when inherited from a source PDF.
_LOCK_CATALOG_KEYS = ("AcroForm", "Perms", "Legal")


def strip_signature_locks(doc) -> tuple:
    """
    Delete signature widgets and null the catalog keys that lock editing.

    Why this is needed (found 2026-08-12 on PO26-1011): ACME Medical drawings released
    through DocuSign carry an AcroForm with /SigFlags 3 — SignaturesExist(1) +
    AppendOnly(2) — plus a signature widget (ENVELOPEID_..., an invisible 0x0 rect).
    PyMuPDF's insert_pdf carries that apparatus into the merged catalog, and
    AppendOnly makes Acrobat refuse page insert/delete/replace and refuse a new
    signature. The package is NOT encrypted and its permission bits look fine, so
    is_encrypted/permissions checks alone will not catch it.

    Returns (removed_widgets, nulled_catalog_keys) where removed_widgets is a list of
    (page_number, field_name, was_visible) so a caller can flag the rare case of a
    VISIBLE signature appearance being dropped.
    """
    removed = []
    for pno in range(doc.page_count):
        page = doc[pno]
        for w in list(page.widgets() or []):
            if w.field_type_string == "Signature":
                removed.append((pno + 1, w.field_name, not w.rect.is_empty))
                page.delete_widget(w)
    cat = doc.pdf_catalog()
    present = doc.xref_object(cat, compressed=False)
    nulled = []
    for key in _LOCK_CATALOG_KEYS:
        if f"/{key}" in present:
            doc.xref_set_key(cat, key, "null")
            nulled.append(key)
    return removed, nulled


def verify_editable(pdf_path: str):
    """
    Prove a compiled package can actually be signed and page-edited.

    Checks the flags AND exercises the real operations (annotate, insert page, delete
    page) on a scratch copy, because the DocuSign AppendOnly lock leaves encryption
    and permission bits looking perfectly healthy. Raises PdfEditabilityError listing
    every problem found; returns None on success.
    """
    problems = []
    doc = fitz.open(pdf_path)
    try:
        cat = doc.xref_object(doc.pdf_catalog(), compressed=False)
        if doc.is_encrypted or doc.needs_pass:
            problems.append("document is encrypted / password protected")
        for key in _LOCK_CATALOG_KEYS:
            if f"/{key}" in cat:
                problems.append(f"catalog still carries /{key}")
        if "SigFlags" in cat:
            problems.append("catalog still carries /SigFlags (AppendOnly blocks editing)")
        stray = [(n + 1, w.field_name) for n in range(doc.page_count)
                 for w in (doc[n].widgets() or []) if w.field_type_string == "Signature"]
        if stray:
            problems.append(f"signature widgets remain: {stray}")
        for bit, label in ((fitz.PDF_PERM_MODIFY, "modify"),
                           (fitz.PDF_PERM_ANNOTATE, "annotate/sign"),
                           (fitz.PDF_PERM_FORM, "fill forms"),
                           (fitz.PDF_PERM_ASSEMBLE, "assemble/reorder pages")):
            if not doc.permissions & bit:
                problems.append(f"permission denied: {label}")
        page_count = doc.page_count
    finally:
        doc.close()

    tmp_dir = TOOLS_DIR / "_tmp"
    tmp_dir.mkdir(exist_ok=True)
    probe = tmp_dir / f"editcheck_{Path(pdf_path).stem}.pdf"
    try:
        doc = fitz.open(pdf_path)
        doc[0].add_freetext_annot(fitz.Rect(40, 40, 200, 70), "signature probe")
        doc.insert_page(-1)
        doc.delete_page(doc.page_count - 1)
        doc.save(str(probe), garbage=4, deflate=True, clean=True)
        doc.close()
        chk = fitz.open(probe)
        if chk.page_count != page_count:
            problems.append(f"page ops changed count: {page_count} -> {chk.page_count}")
        if not (chk[0].annots() or []):
            problems.append("annotation did not persist (cannot sign)")
        chk.close()
    except Exception as e:
        problems.append(f"functional edit test failed: {type(e).__name__}: {e}")
    finally:
        try:
            probe.unlink(missing_ok=True)
        except OSError:
            pass

    if problems:
        raise PdfEditabilityError(
            f"{Path(pdf_path).name} is NOT fully editable — refusing to deliver a "
            f"package that cannot be signed:\n  " + "\n  ".join(problems))


class UserContentError(Exception):
    """
    Raised when regenerating a compiled package would destroy content a human added
    to the previously delivered copy (Section 4 entries, signatures, review markup).
    Never catch and ignore this — overwriting a signed quality record is unrecoverable.
    """
    pass


def find_user_edits(delivered_path: str, fresh_path: str) -> dict:
    """
    Detect content a human added to an already-delivered compiled package.

    Compares the delivered file against a freshly regenerated one and reports, per
    page, text tokens and annotation types present in the delivered copy but absent
    from the regenerated one. Token-set comparison (not a pixel or byte diff) is used
    deliberately: Acrobat re-encodes and re-renders pages when it saves, so bytes and
    even rendered pixels change without any content being added. Verified on
    PO26-1011 — the PRF Sections 1-2 page rendered differently after Acrobat saved it
    while its text tokens were identical.

    Returns {page_number: {"text": [...], "annots": [...]}} for pages with additions;
    an empty dict means nothing was added.
    """
    edits = {}
    a = fitz.open(delivered_path)
    b = fitz.open(fresh_path)
    try:
        for pno in range(min(a.page_count, b.page_count)):
            added_text = sorted(set((a[pno].get_text() or "").split()) -
                                set((b[pno].get_text() or "").split()))
            fresh_annots = [an.type[1] for an in (b[pno].annots() or [])]
            added_annots = []
            for an in (a[pno].annots() or []):
                kind = an.type[1]
                if kind in fresh_annots:
                    fresh_annots.remove(kind)
                else:
                    added_annots.append(kind)
            if added_text or added_annots:
                edits[pno + 1] = {"text": added_text, "annots": added_annots}
    finally:
        a.close()
        b.close()
    return edits


def assert_no_user_edits(delivered_path: str, fresh_path: str, force: bool = False):
    """
    Refuse to overwrite a delivered package that a human has since written on.

    Call this immediately before replacing an existing compiled package. Pass
    force=True only on an explicit operator instruction to discard their edits.

    Why this exists (PO26-1011, 2026-08-12): the engineer types the lot release report
    names and the accept/reject decision straight into Section 4 of the COMPILED PRF,
    and the Quality Manager signs the compiled package. Re-running a PO script to pick
    up a corrected supplier document would silently overwrite all of it. On that order
    a re-run was stopped only by luck — Acrobat happened to hold a file lock on the
    one package that already had Section 4 filled in.
    """
    if force or not Path(delivered_path).exists():
        return
    edits = find_user_edits(delivered_path, fresh_path)
    if not edits:
        return
    lines = []
    for pno, info in sorted(edits.items()):
        if info["text"]:
            preview = " ".join(info["text"])[:220]
            lines.append(f"page {pno}: added text -> {preview}")
        if info["annots"]:
            lines.append(f"page {pno}: added annotations -> {info['annots']}")
    raise UserContentError(
        f"{Path(delivered_path).name} has content added since it was generated — "
        f"refusing to overwrite it:\n  " + "\n  ".join(lines) +
        "\n  Move or rename the edited copy first, or re-run with force to discard it.")


def compile_pdf(doc_paths: list, output_path: str, flatten_paths: list = None):
    """
    Merge PDFs in order into a fully editable, signable compiled package.

    1. PyMuPDF merges all source PDFs in the specified order into a temp file.
    2. strip_signature_locks() removes signature widgets and the /AcroForm,
       /Perms, /Legal catalog locks inherited from source PDFs (DocuSign-released
       drawings carry /SigFlags AppendOnly, which blocks signing and page edits).
    3. Ghostscript re-distills → strips any remaining security/signature/encryption
       metadata and flattens the result.
    4. verify_editable() gates the output: raises PdfEditabilityError rather than
       hand back a package the Quality Manager cannot sign.

    The count of ☑ (U+2611) text glyphs is compared before and after re-distill and a
    change raises. See the WI-7.4.3-01 Step 3 CAUTION: Ghostscript /prepress mangles
    NATIVE form-control checkbox appearances into X's. Scripts that render the marks
    as ☑/☐ text glyphs are unaffected (measured 13 before / 13 after on PO26-1011),
    but if a PRF relies on native form controls, run normalize_checkbox_marks() on the
    output afterwards to turn the X's back into checkmarks.

    Temp files are written to TOOLS_DIR/_tmp (never inside a PO folder).
    If Ghostscript is not installed, falls back to the PyMuPDF strip with a warning —
    the editability gate still applies.
    Install Ghostscript: https://ghostscript.com/releases/gsdnld.html
    """
    tmp_dir = TOOLS_DIR / "_tmp"
    tmp_dir.mkdir(exist_ok=True)
    tmp_merged = str(tmp_dir / "merged_raw.pdf")

    merged = fitz.open()
    for path in doc_paths:
        doc = fitz.open(path)
        merged.insert_pdf(doc)
        doc.close()

    removed, nulled = strip_signature_locks(merged)
    for pno, name, visible in removed:
        note = ("VISIBLE appearance dropped — check that page"
                if visible else "invisible, no content lost")
        print(f"  Stripped signature widget p{pno}: {name} ({note})")
    if nulled:
        print(f"  Nulled catalog lock(s): {', '.join('/' + k for k in nulled)}")

    marks_before = sum(p.get_text().count("☑") for p in merged)
    merged.save(tmp_merged, garbage=4, deflate=True, clean=True)
    merged.close()

    success = _gs_redistill(tmp_merged, output_path)
    if not success:
        shutil.copy(tmp_merged, output_path)
        print(
            "WARNING: Ghostscript not found — relying on the PyMuPDF lock strip alone.\n"
            "Install from https://ghostscript.com/releases/gsdnld.html\n"
            "Then re-run to get a fully re-distilled compiled PDF."
        )

    with fitz.open(output_path) as out:
        marks_after = sum(p.get_text().count("☑") for p in out)
    if marks_after != marks_before:
        raise ChecklistCheckboxError(
            f"Checkmark count changed during compilation of "
            f"{Path(output_path).name}: {marks_before} -> {marks_after}. This is the "
            f"WI-7.4.3-01 Step 3 CAUTION case — refusing to deliver a quality record "
            f"whose checkboxes may render as X's.")

    verify_editable(output_path)


_LOT_SUBROW_RE = re.compile(r"\blot\s*#?\s*:", re.IGNORECASE)


def _is_lot_subrow(text: str) -> bool:
    """
    True if a text line looks like the Lot#/ExpDate/Rev#/Mft Date/Ster Date
    sub-row of a receiving-list row (e.g. "Lot#: 9900001 ExpDate: ..."). Used
    as a row-boundary marker: on these lists a sub-row is the LAST line that
    belongs to a given product entry — the very next line, if any, is always
    the NEXT product's own main line. Detecting this line by content (rather
    than by vertical gap alone) is what lets highlight_receiving_line() stop
    at the correct row boundary even on continuous, evenly-spaced lists where
    there's no gap to detect (see Open Item #5 in CLAUDE.md).
    """
    return bool(_LOT_SUBROW_RE.search(text or ""))


# Receiving-date / quantity detectors used to pick the ONE row a highlight
# belongs on when the search term appears on several rows of the same list.
# Added 2026-09-08 (found on PO26-1007/VendorE — a 4-shipment blanket order whose
# receiving list carries one row per scheduled shipment, all with the SAME
# part number and no Lot# column at all). See highlight_receiving_line().
_RECV_DATE_RE = re.compile(r"\b\d{1,2}/\d{1,2}/\d{2,4}\b")
_RECV_QTY_RE  = re.compile(r"\b\d[\d,]*\s*(?:ea|each|pcs?|rolls?|units?|bx|box(?:es)?)\b",
                           re.IGNORECASE)


def _has_receiving_date(text: str) -> bool:
    """True if a receiving-list line carries a receiving date (M/D/YY[YY])."""
    return bool(_RECV_DATE_RE.search(text or ""))


def _has_received_qty(text: str) -> bool:
    """True if a receiving-list line carries a quantity cell (e.g. '10 ea')."""
    return bool(_RECV_QTY_RE.search(text or ""))


def _is_received_row(text: str) -> bool:
    """
    True if a receiving-list line is an ACTUALLY-RECEIVED row: it has both a
    receiving date and a quantity. On a blanket/scheduled order the list
    contains one row per scheduled shipment — only the shipment(s) already
    received carry a receiving date and a Qty Received value; future rows
    carry the part number and Qty To Receive only. The highlight must land on
    the received row, not on whichever row the part-number search happened to
    hit first.
    """
    return _has_receiving_date(text) and _has_received_qty(text)


# Sterilization-service round-trip disambiguation — added 2026-09-11 (WI-
# 7.4.3-01 gap analysis, PO26-1005/PO26-1009, SterilizerCo). On a
# sterilization-vendor receiving list the SAME lot number legitimately
# appears on two different rows: an unsterile-input entry (lot sent out) and
# a sterile-output entry (lot received back) — both can carry a receiving
# date and a quantity, so _is_received_row() alone cannot tell them apart
# the way it disambiguates a blanket order's future/past shipment rows. The
# row this receiving inspection actually cares about is the STERILE
# RETURN — that's the lot physically being received and inspected today.
_STERILE_OUTPUT_RE = re.compile(
    r"\b(sterile|sterili[sz]ed?|post[- ]?ster|irradiat\w*|gamma|e-?beam|dose)\b",
    re.IGNORECASE)
_RECV_DATE_PARSE_RE = re.compile(r"\b(\d{1,2})/(\d{1,2})/(\d{2,4})\b")


def _is_sterile_output_row(text: str) -> bool:
    """True if a receiving-list line's own text names it as the sterile/
    post-sterilization side of a round trip (vs. the raw/unsterile input)."""
    return bool(_STERILE_OUTPUT_RE.search(text or ""))


def _parse_recv_date(text: str):
    """Best-effort M/D/YY(YY) -> a sortable (year, month, day) tuple, or None
    if no date is present or it doesn't parse. Used only to break a tie
    between several equally-qualified candidate rows — never to validate a
    date's correctness."""
    m = _RECV_DATE_RE.search(text or "")
    if not m:
        return None
    m2 = _RECV_DATE_PARSE_RE.search(text)
    if not m2:
        return None
    mo, da, yr = m2.groups()
    try:
        mo, da, yr = int(mo), int(da), int(yr)
        if yr < 100:
            yr += 2000
        return (yr, mo, da)
    except ValueError:
        return None


def _bold_line_flags(page, line_rects) -> list:
    """
    Return [bool] parallel to line_rects — True where that grouped text line
    contains bold text. Used as a hard row-boundary: a receiving list's column
    heading ("Part | Receiving Date | Receiving Location | Qty To Receive |
    Qty Received") and its page title/order-info block are bold, and the
    highlight must never cover heading text (explicit QE instruction,
    2026-09-08). Matching is done by span-rect overlap rather than by block/
    line index so it does not depend on "words" and "dict" numbering agreeing.

    Bold is detected from the span's own render flags (PyMuPDF bit 4 = bold)
    with a font-name fallback, since some generators embed a bold face without
    setting the flag.
    """
    bold_rects = []
    try:
        page_dict = page.get_text("dict")
    except Exception:
        return [False] * len(line_rects)
    for blk in page_dict.get("blocks", []):
        for ln in blk.get("lines", []):
            for sp in ln.get("spans", []):
                if not (sp.get("text") or "").strip():
                    continue
                fname = (sp.get("font") or "").lower()
                if (sp.get("flags", 0) & 16) or any(
                        k in fname for k in ("bold", "black", "heavy", "semibold")):
                    bold_rects.append(fitz.Rect(sp["bbox"]))
    flags = []
    for r in line_rects:
        hit = False
        for b in bold_rects:
            b_cy = (b.y0 + b.y1) / 2
            if r.y0 <= b_cy <= r.y1 and b.x1 > r.x0 and b.x0 < r.x1:
                hit = True
                break
        flags.append(hit)
    return flags


def highlight_receiving_line(receiving_pdf_path: str, search_term: str, output_path: str,
                              max_row_span: float = 60.0, line_gap_factor: float = 1.8,
                              part_number: str = "") -> bool:
    """
    Produce a copy of the Receiving List with the FULL multi-line row for
    search_term highlighted in yellow — not just the single text line the
    search term happens to appear on, and not bleeding into a neighboring
    product's row. Use this copy (not the original) when compiling.

    Root cause fixed 2026-07-20: this previously highlighted only a single
    line (inst.y0-2 to inst.y1+2 around the search hit). On this vendor's
    receiving list layout, a row is TWO lines — a main line (part name /
    description / qty) and a sub-row directly below it (Lot#, ExpDate, Rev#,
    Mft Date, Ster Date). Searching on the part number (which lands on the
    main line) highlighted only Part Number + Quantity; the Lot/Date/Rev
    sub-row was left unhighlighted. Per WI-7.4.3-01 Step 2, the highlight
    must cover the full row, part name through the lot# sub-row.

    search_term should be the LOT NUMBER, not the part number (see
    WI-7.4.3-01 Step 2 "HIGHLIGHT SEARCH TERM" rule) — the caller is
    responsible for passing the lot number once it has been extracted, and
    falling back to the part number only if lot extraction failed.

    Row-expansion heuristic (revised 2026-07-20, Open Item #5): starting from
    the text line containing search_term, walks upward and downward through
    the page's text lines (grouped by PyMuPDF's own block/line numbering).
    Two independent stopping conditions both have to hold for a neighboring
    line to be pulled in:
      1. Vertical gap test (original heuristic) — the gap to the neighboring
         line is no more than line_gap_factor times the page's typical line
         spacing, and the total span stays under max_row_span.
      2. Content boundary test (new) — expansion never crosses a line that
         looks like a Lot#/sub-row line (see `_is_lot_subrow`) OTHER than the
         current row's own sub-row: walking upward stops before absorbing the
         PREVIOUS row's sub-row, and walking downward stops as soon as this
         row's own sub-row has been included (it never continues on into the
         NEXT row's main line).
    This was added because on a receiving list with no blank-line separator
    between products (uniform line spacing throughout the page), condition 1
    alone never triggers a stop — it will walk straight through into the next
    product's lines too (observed on PO26-1002's VendorA receiving list: both
    SR90-1008's and SR90-1009's rows got highlighted together, including both
    quantity cells, instead of just the one row that matched the search term).

    The highlight rect is now the union of the included lines' own text
    bounding boxes (with a small pad), not a page-margin-to-page-margin box —
    so it hugs the row's actual text instead of highlighting blank space and
    making the over-highlight look worse than it is.

    Returns True if search_term was found and at least one row was
    highlighted, False if it was not found anywhere in the PDF's text layer
    (scanned page / no text layer) — callers should flag a False return
    rather than silently including an unhighlighted page (per WI-7.4.3-01
    Step 2).

    part_number (added 2026-09-11, corrected same day — see below): an
    optional secondary anchor, used ONLY to disambiguate a specific
    sterilization-round-trip case where the documented lot-number anchor
    (search_term) matches more than one row and NONE of them independently
    qualifies as a received row.

    CORRECTED 2026-09-11: this session's first attempt at this fix added a
    sterile/irradiated KEYWORD test on the lot-matched rows themselves. Real
    per-PO scripts proven against the actual SterilizerCo receiving lists
    (process_PO26-1005.py / process_PO26-1009.py, confirmed via the device
    bridge) show that is not what the real, working fix does, and the
    keyword approach does not even engage on the documented failure: on
    PO26-1009's real list, the unsterile SR90-1010 row and the sterile
    SR90-1011 row share the SAME Lot#: 0363/1, and BOTH hits from a
    lot-number search land on Lot#/sub-row lines — neither qualifies as
    `_is_received_row()` at all, so `received` is EMPTY, not "more than
    one" (the case this function's first pass actually handled). The real,
    verified fix the bespoke script uses is to switch the anchor entirely:
    when the PART NUMBER is unique on the page, search on THAT instead of
    the lot number — it lands directly on the correct row's own main line
    (verified: PO26-1009's SR90-1011 row, not the SR90-1010 row a bare
    lot-number search falls back to). This function now does that retry
    internally when part_number is supplied: if the lot-number search
    (search_term) yields 2+ candidate rows and NONE qualifies as received,
    and part_number is provided, non-empty, and different from search_term,
    it re-searches THIS SAME PAGE for part_number and — if that search lands
    on exactly one row — highlights that row instead. The sterile-keyword/
    latest-date fallback below is kept as a secondary safety net for when
    part_number isn't supplied, doesn't disambiguate either, or the
    ambiguity takes the ORIGINAL "more than one received row" shape instead
    (which the keyword logic below was actually verified capable of
    resolving, just not the shared-Lot#-with-zero-received-rows case).
    """
    doc = fitz.open(receiving_pdf_path)
    any_match = False
    for page in doc:
        hits = page.search_for(search_term)
        if not hits:
            continue
        any_match = True

        words = page.get_text("words")  # (x0, y0, x1, y1, word, block_no, line_no, word_no)
        if not words:
            continue

        # Group words into lines using PyMuPDF's own (block_no, line_no) keys —
        # more reliable than re-clustering by raw y-coordinate.
        line_groups = {}
        for w in words:
            x0, y0, x1, y1, word = w[0], w[1], w[2], w[3], w[4]
            key = (w[5], w[6])
            line_groups.setdefault(key, []).append((x0, y0, x1, y1, word))

        line_rects, line_texts = [], []
        for boxes in line_groups.values():
            lx0 = min(b[0] for b in boxes); ly0 = min(b[1] for b in boxes)
            lx1 = max(b[2] for b in boxes); ly1 = max(b[3] for b in boxes)
            line_rects.append(fitz.Rect(lx0, ly0, lx1, ly1))
            # Reconstruct the line's text in left-to-right order for the
            # content boundary test (_is_lot_subrow).
            ordered = sorted(boxes, key=lambda b: b[0])
            line_texts.append(" ".join(b[4] for b in ordered))
        order = sorted(range(len(line_rects)), key=lambda i: line_rects[i].y0)
        line_rects = [line_rects[i] for i in order]
        line_texts = [line_texts[i] for i in order]

        gaps = [line_rects[i + 1].y0 - line_rects[i].y1
                for i in range(len(line_rects) - 1) if line_rects[i + 1].y0 > line_rects[i].y1]
        # Exclude extreme outlier gaps (e.g. the gap down to a page footer)
        # before taking the median — found 2026-07-27 on PO26-1006/VendorK: a
        # short receiving list with only two line items has few gaps to
        # sample, and the huge footer gap (500+pt) was large enough to pull
        # the median itself well above the real ~24pt boundary gap between
        # the two products' rows. That made the ordinary row-boundary gap
        # look insignificant by comparison, so the gap test never fired and
        # the highlight walked straight through into the neighboring
        # product's row. Capping which gaps count toward the median at
        # 2x max_row_span keeps the statistic representative of normal
        # line/row spacing on the page, not skewed by a single non-content
        # outlier — a table with only single-line items (no row ever
        # legitimately spans more than max_row_span) will never have a real
        # boundary gap that large anyway.
        typical_gaps = [g for g in gaps if g <= max_row_span * 2] or gaps
        median_gap = sorted(typical_gaps)[len(typical_gaps) // 2] if typical_gaps else 2.0

        is_bold_line = _bold_line_flags(page, line_rects)

        # ── Group the page's text lines into visual ROWS ─────────────────────
        # 2026-09-08 (found on PO26-1007/VendorE): grouping words by PyMuPDF's own
        # (block_no, line_no) does NOT yield one entry per visual row on a
        # table-shaped receiving list — it yields one entry per CELL. On this
        # vendor's list each row's Part, Receiving Date, Qty To Receive and Qty
        # Received cells are four separate blocks sharing the same y-band, so
        # the previous "walk up and down through neighbouring lines" model was
        # walking through cells that sit side by side, not through the rows
        # above and below. That is why the row containing the search hit
        # appeared to have no receiving date on it (the date is a sibling cell,
        # not part of that text line) and why the walk climbed into the bold
        # column heading.
        #
        # Cluster by vertical-CENTRE proximity, not by extent overlap. Cells
        # that share a visual row have near-identical centres (they are set on
        # the same baseline), whereas lines stacked one above the other differ
        # by the line pitch — even when their bounding boxes overlap slightly,
        # which they routinely do on tightly-spaced lists. An overlap-based
        # test merges those stacked lines into one "row" and defeats the whole
        # point of the grouping (verified against a two-line Part/Lot#-subrow
        # fixture, where an overlap test merged a product's sub-row with the
        # NEXT product's main line). Tolerance is half the shorter line's
        # height, so same-baseline cells always group and separate baselines
        # never do.
        #
        # On a single-column list (one text line per row) every row has
        # exactly one line, so this reduces to the previous behaviour; on the
        # two-line Part/Lot#-subrow layout the main line and its sub-row
        # remain two separate rows and the downward walk still absorbs the
        # sub-row exactly as before.
        rows = []
        for i in range(len(line_rects)):
            r = line_rects[i]
            r_cy = (r.y0 + r.y1) / 2
            for row in rows:
                rr = row["rect"]
                rr_cy = (rr.y0 + rr.y1) / 2
                if abs(r_cy - rr_cy) <= 0.5 * min(r.height, rr.height):
                    row["idxs"].append(i)
                    row["rect"] = fitz.Rect(min(rr.x0, r.x0), min(rr.y0, r.y0),
                                            max(rr.x1, r.x1), max(rr.y1, r.y1))
                    break
            else:
                rows.append({"idxs": [i], "rect": fitz.Rect(r)})
        rows.sort(key=lambda row: row["rect"].y0)

        row_rects = [row["rect"] for row in rows]
        row_texts = [" ".join(line_texts[i] for i in
                              sorted(row["idxs"], key=lambda j: line_rects[j].x0))
                     for row in rows]
        row_bold  = [any(is_bold_line[i] for i in row["idxs"]) for row in rows]

        # Recompute the typical gap from ROW boundaries. Computing it from the
        # per-cell line rects (as this used to) mixes in the ~0pt "gaps"
        # between cells of the same row and drags the median far below real
        # row spacing, which makes the gap test fire on every neighbour.
        row_gaps = [row_rects[i + 1].y0 - row_rects[i].y1
                    for i in range(len(row_rects) - 1)
                    if row_rects[i + 1].y0 > row_rects[i].y1]
        typical_row_gaps = [g for g in row_gaps if g <= max_row_span * 2] or row_gaps
        median_gap = (sorted(typical_row_gaps)[len(typical_row_gaps) // 2]
                      if typical_row_gaps else 2.0)

        # ── Map every hit to its own visual row ──────────────────────────────
        # Match the hit to its own line by CENTER containment, not a fuzzy
        # tolerance on y0. A fixed ±1pt tolerance on inst.y0 alone can snap
        # a hit onto the wrong (adjacent) line when lines are packed tight
        # (e.g. a hit whose y0 sits within 1pt of the PRECEDING line's y1)
        # — found 2026-07-20 while testing Open Item #5's fix: searching a
        # part number on a densely-packed list incorrectly matched the
        # previous row's own sub-row line instead of this row's main line,
        # which made the whole row-expansion walk start from the wrong
        # line entirely. The hit's own vertical center is always well
        # inside its true line's bounding box (that box is the union of
        # every word on the line, which includes the hit) and essentially
        # never inside a neighboring line's box, so this is unambiguous.
        hit_line_idxs, unmapped_hits = [], []
        for inst in hits:
            hit_cy = (inst.y0 + inst.y1) / 2
            containing = [i for i, r in enumerate(row_rects) if r.y0 <= hit_cy <= r.y1]
            if containing:
                hit_line_idxs.append(containing[0])
            elif row_rects:
                hit_line_idxs.append(min(
                    range(len(row_rects)),
                    key=lambda i: abs((row_rects[i].y0 + row_rects[i].y1) / 2 - hit_cy)))
            else:
                unmapped_hits.append(inst)

        for inst in unmapped_hits:
            # Couldn't map the hit to a grouped line — fall back to the
            # old single-line behavior rather than skip highlighting entirely.
            line_rect = fitz.Rect(inst.x0 - 2, inst.y0 - 2, inst.x1 + 2, inst.y1 + 2)
            annot = page.add_highlight_annot(line_rect)
            annot.set_colors(stroke=(1, 1, 0))
            annot.update()

        # ── Pick the ONE row to highlight ───────────────────────────────────
        # 2026-09-08 fix (found on PO26-1007/VendorE, reported by QE on the first
        # compiled package): this loop used to highlight EVERY search hit, one
        # annotation each. That is correct only when the search term is unique
        # on the page — which is what the documented lot-number anchor
        # guarantees. On a blanket/scheduled order the receiving list carries
        # one row per scheduled shipment, all with the SAME part number and
        # (on this vendor's layout) no Lot# column at all, so lot extraction
        # returns nothing and process_part() falls back to the part number.
        # That matched all four shipment rows and highlighted all four.
        # Per explicit QE instruction: highlight only the first row that has
        # BOTH a receiving date and a received quantity — i.e. the shipment
        # actually being received — and never any heading text.
        # When the search term matches exactly one row — which is what the
        # documented lot-number anchor normally guarantees — that row IS the
        # target and no selection is needed. The received-row preference below
        # only engages when the term is ambiguous across several rows, so this
        # cannot change the outcome for any list that was already correct.
        candidates = sorted(set(hit_line_idxs))
        if len(candidates) <= 1:
            targets = candidates
        else:
            received = [i for i in candidates if _is_received_row(row_texts[i])]
            dated    = [i for i in candidates if _has_receiving_date(row_texts[i])]

            # ── Part-number-anchor retry (2026-09-11, corrected same day) ────
            # Real failure mode confirmed against PO26-1005/PO26-1009's actual
            # receiving lists: the lot number is shared by an unsterile-input
            # row and a sterile-output row, and BOTH hits land on Lot#/sub-row
            # lines — `received` comes back EMPTY, not "more than one" (see
            # this function's docstring). Only engage this retry in exactly
            # that shape (no received row at all, 2+ lot-search candidates,
            # a usable part_number given) so it can never change behavior for
            # a list that already resolves correctly via `received`/`dated`.
            if not received and part_number and part_number != search_term:
                pn_hits = page.search_for(part_number)
                pn_line_idxs = []
                for inst in pn_hits:
                    hit_cy = (inst.y0 + inst.y1) / 2
                    containing = [i for i, r in enumerate(row_rects) if r.y0 <= hit_cy <= r.y1]
                    if containing:
                        pn_line_idxs.append(containing[0])
                pn_candidates = sorted(set(pn_line_idxs))
                if len(pn_candidates) == 1:
                    targets = pn_candidates
                    print(f"  Receiving list: '{search_term}' (lot#) matched "
                          f"{len(candidates)} rows and none qualified as received "
                          f"(sterilization round trip — same lot# on an unsterile-"
                          f"input row and a sterile-output row) — re-anchored on "
                          f"part number '{part_number}', which uniquely identified "
                          f"the correct row.")
                    received = None  # sentinel: skip the received/dated branches below
            if received is None:
                pass
            elif received:
                if len(received) == 1:
                    targets = [received[0]]
                    print(f"  Receiving list: '{search_term}' matched {len(candidates)} rows "
                          f"(scheduled/blanket order) — highlighting the received row only.")
                else:
                    # 2026-09-11 fix (found on PO26-1005/PO26-1009, SterilizerCo):
                    # more than one row can independently qualify as "received"
                    # (has both a date and a qty) — the blanket-order case
                    # above assumed at most one such row, which does not hold
                    # on a sterilization-service round trip: the SAME lot#
                    # legitimately labels both an unsterile-input row and a
                    # sterile-output row, and both can carry a receiving date
                    # and quantity. Picking "the first" arbitrarily risked
                    # highlighting the wrong leg of the round trip. Prefer a
                    # row whose own text names it as the sterile/post-
                    # sterilization side — that is the lot actually being
                    # received and inspected today. If none/more than one
                    # matches that keyword, fall back to the row with the
                    # latest parseable receiving date (the round trip's
                    # later, "coming back" event); if dates don't
                    # discriminate either, keep the first candidate but warn
                    # loudly so a human verifies by eye rather than silently
                    # trusting a guess.
                    sterile_out = [i for i in received if _is_sterile_output_row(row_texts[i])]
                    if len(sterile_out) == 1:
                        targets = sterile_out
                        print(f"  Receiving list: '{search_term}' matched {len(received)} "
                              f"received rows (sterilization round trip) — highlighting the "
                              f"sterile/post-sterilization row.")
                    else:
                        pool = sterile_out if sterile_out else received
                        dated_pairs = [(i, _parse_recv_date(row_texts[i])) for i in pool]
                        dated_pairs = [(i, d) for i, d in dated_pairs if d is not None]
                        if dated_pairs:
                            targets = [max(dated_pairs, key=lambda p: p[1])[0]]
                        else:
                            targets = [pool[0]]
                        print(f"  [WARN] highlight_receiving_line: '{search_term}' matched "
                              f"{len(received)} received rows and the sterile/output row "
                              f"could not be uniquely identified — highlighting the "
                              f"latest-dated candidate row. This may be the unsterile input "
                              f"row, not the sterile return — verify by eye before compiling.")
            elif dated:
                targets = [dated[0]]
                print(f"  [WARN] highlight_receiving_line: '{search_term}' matched "
                      f"{len(candidates)} rows and none has both a receiving date and a "
                      f"quantity — highlighting the first dated row. Verify by eye.")
            else:
                targets = [candidates[0]]
                print(f"  [WARN] highlight_receiving_line: '{search_term}' matched "
                      f"{len(candidates)} rows and none carries a receiving date — "
                      f"highlighting the first matching row only. Verify by eye.")

        for idx in targets:
            top = bottom = idx
            # Expand upward: stop at the vertical-gap boundary, never absorb
            # another sub-row line (that would be the PREVIOUS row's own
            # Lot#/date/rev line, not this row's), and never absorb a BOLD
            # line — on a receiving list the bold lines are the column heading
            # and the page title/order-info block, which must never be
            # highlighted (explicit QE instruction 2026-09-08; the old code
            # walked straight up into the "Part | Receiving Date | ... | Qty
            # Received" heading on this layout, because _is_lot_subrow never
            # matches a list with no Lot# column and the line spacing is
            # uniform right through the heading).
            while top > 0:
                candidate = row_rects[top - 1]
                gap = row_rects[top].y0 - candidate.y1
                span = row_rects[idx].y1 - candidate.y0
                if (_is_lot_subrow(row_texts[top - 1]) or row_bold[top - 1]
                        or (top - 1) in candidates):
                    break
                if gap <= median_gap * line_gap_factor and span <= max_row_span:
                    top -= 1
                    # 2026-09-17 fix (found on PO26-1004/VendorB — the Open Item #5
                    # bleed class on a layout none of its earlier fixes cover).
                    # Stop as soon as the line just absorbed is itself a
                    # received row, i.e. this row's own main line, the one
                    # carrying the receiving date and the quantity cells.
                    # Nothing ABOVE a row's main line can belong to that row,
                    # so continuing past it can only walk into the row above.
                    # On PO26-1004 the anchor is the lot number (the documented
                    # anchor), which sits on the Lot#/Rev#/Mft-Date sub-row, so
                    # the walk correctly climbs to the SR90-1002 main line — and
                    # then kept going into the WRAPPED SECOND LINE of the
                    # "Manufacturing Costs" row above, which is neither bold,
                    # nor a Lot# sub-row, nor another candidate, so none of the
                    # three break tests above caught it. This is the exact
                    # mirror of the downward walk's own _is_received_row(nxt)
                    # break below, and it can only ever SHRINK a highlight,
                    # never extend one.
                    if _is_received_row(row_texts[top]):
                        break
                else:
                    break
            # Expand downward: stop at the vertical-gap boundary, never
            # continue past this row's own sub-row line into the NEXT row's
            # main line, never cross a bold line, and never absorb another
            # received/dated row (that is a different shipment's row).
            already_have_subrow = any(_is_lot_subrow(t) for t in row_texts[idx:bottom + 1])
            while bottom < len(row_rects) - 1 and not already_have_subrow:
                candidate = row_rects[bottom + 1]
                gap = candidate.y0 - row_rects[bottom].y1
                span = candidate.y1 - row_rects[idx].y0
                nxt = row_texts[bottom + 1]
                if (row_bold[bottom + 1] or _is_received_row(nxt)
                        or (bottom + 1) in candidates):
                    break
                if gap <= median_gap * line_gap_factor and span <= max_row_span:
                    bottom += 1
                    if _is_lot_subrow(row_texts[bottom]):
                        already_have_subrow = True
                else:
                    break

            x0 = min(row_rects[i].x0 for i in range(top, bottom + 1))
            x1 = max(row_rects[i].x1 for i in range(top, bottom + 1))
            # Clamp the vertical pad to at most half the gap to the adjacent
            # row, so the highlight can never visually bleed onto the row above
            # or below on a tightly-spaced list (2026-09-08: a flat 2pt pad
            # overlapped the neighbouring row on a layout whose rows sit
            # 0.7pt apart).
            pad_top = 2.0 if top == 0 else min(
                2.0, max(0.0, (row_rects[top].y0 - row_rects[top - 1].y1) / 2))
            pad_bot = 2.0 if bottom >= len(row_rects) - 1 else min(
                2.0, max(0.0, (row_rects[bottom + 1].y0 - row_rects[bottom].y1) / 2))
            row_rect = fitz.Rect(x0 - 2, row_rects[top].y0 - pad_top,
                                 x1 + 2, row_rects[bottom].y1 + pad_bot)
            annot = page.add_highlight_annot(row_rect)
            annot.set_colors(stroke=(1, 1, 0))
            annot.update()

    doc.save(output_path)
    doc.close()
    return any_match


# ─────────────────────────────────────────────
# CHECKBOX MARK NORMALIZATION — X → CHECKMARK
# ─────────────────────────────────────────────

_PRF_PAGE_MARKERS = ("INSPECTION CHECKLIST", "PART INFORMATION",
                     "DISCREPANCY DISPOSITION", "DATA REPORT", "COMMENTS")


def count_checkbox_appearances(pdf_path: str) -> tuple:
    """
    Return (boxes, marked) for the PRF pages of an exported PDF: how many
    checkbox border squares were rendered, and how many of those contain
    diagonal (X) strokes — i.e. how many are CHECKED as LibreOffice draws a
    checked native form control.

    Added 2026-09-17. Since the ☑/☐ text-glyph overlay was removed from
    fill_prf() (QE's instruction: the template's own native checkboxes are
    the checkboxes of record), the native control's rendered appearance is the
    ONLY thing carrying an answer onto the signed PDF. That is safe on this
    LibreOffice build — verified by rendering — but the long-standing claim in
    this codebase was the opposite ("LibreOffice cannot render native VML
    checkboxes at all"), so a different build or version could plausibly
    behave that way and every checkbox would silently export BLANK. This
    counter exists so process_part() can detect exactly that and refuse to
    deliver the package, instead of shipping a quality record whose answers
    are invisible.

    Same geometric detection as normalize_checkbox_marks() below, which is
    what runs immediately afterwards — deliberately so: if this function can
    see the boxes and the marks, so can that one.
    """
    doc = fitz.open(pdf_path)
    boxes_total = marked_total = 0
    try:
        for pg in doc:
            if not any(m in pg.get_text() for m in _PRF_PAGE_MARKERS):
                continue
            boxes, diag_mids = [], []
            for d in pg.get_drawings():
                for it in d["items"]:
                    if it[0] == "re":
                        r = it[1]
                        if (4 <= r.width <= 9 and 4 <= r.height <= 9
                                and abs(r.width - r.height) < 2 and r.y0 > 55):
                            boxes.append(fitz.Rect(r))
                    elif it[0] == "l":
                        p, q = it[1], it[2]
                        if abs(q.x - p.x) > 1.2 and abs(q.y - p.y) > 1.2:
                            diag_mids.append(((p.x + q.x) / 2, (p.y + q.y) / 2))
            uniq = []
            for b in boxes:
                c = ((b.x0 + b.x1) / 2, (b.y0 + b.y1) / 2)
                if not any(abs(c[0] - (u.x0 + u.x1) / 2) < 2
                           and abs(c[1] - (u.y0 + u.y1) / 2) < 2 for u in uniq):
                    uniq.append(b)
            boxes_total += len(uniq)
            marked_total += sum(
                1 for b in uniq
                if any(b.x0 <= mx <= b.x1 and b.y0 <= my <= b.y1 for mx, my in diag_mids))
    finally:
        doc.close()
    return boxes_total, marked_total


def normalize_checkbox_marks(pdf_path: str) -> int:
    """
    Force every CHECKED PRF checkbox to render as a clean checkmark (✓), not an X.

    Why this exists: LibreOffice renders native Excel form-control checkboxes
    non-deterministically — sometimes as a checkmark bitmap, sometimes as a vector
    "X" — and the Ghostscript /prepress re-distill further mangles the checkmark
    glyph into an X. Engineers require checkmarks on the signed record, so this
    normalizes the appearance directly on the exported PRF PDF:

      * finds checkbox border rects (~5-9 pt squares) on PRF pages only,
      * a box is "checked" if it contains diagonal (X) strokes,
      * white-outs the X interior (keeping the border) and draws a vector checkmark.

    Boxes that already rendered as checkmark bitmaps have no diagonal strokes and are
    left untouched. Run this on the stamped PRF PDF BEFORE compiling, and compile with
    a PyMuPDF clean merge (NOT Ghostscript) so the vector checkmarks survive.
    Returns the number of checkmarks drawn. Edits the file in place (incremental save).
    """
    doc = fitz.open(pdf_path)
    total = 0
    for pg in doc:
        if not any(m in pg.get_text() for m in _PRF_PAGE_MARKERS):
            continue
        boxes, diag_mids = [], []
        for d in pg.get_drawings():
            for it in d["items"]:
                if it[0] == "re":
                    r = it[1]
                    if (4 <= r.width <= 9 and 4 <= r.height <= 9
                            and abs(r.width - r.height) < 2 and r.y0 > 55):
                        boxes.append(fitz.Rect(r))
                elif it[0] == "l":
                    p, q = it[1], it[2]
                    if abs(q.x - p.x) > 1.2 and abs(q.y - p.y) > 1.2:
                        diag_mids.append(((p.x + q.x) / 2, (p.y + q.y) / 2))
        uniq = []   # dedupe nested border rects by center
        for b in boxes:
            c = ((b.x0 + b.x1) / 2, (b.y0 + b.y1) / 2)
            if not any(abs(c[0] - (u.x0 + u.x1) / 2) < 2 and abs(c[1] - (u.y0 + u.y1) / 2) < 2
                       for u in uniq):
                uniq.append(b)
        for b in uniq:
            if not any(b.x0 <= mx <= b.x1 and b.y0 <= my <= b.y1 for mx, my in diag_mids):
                continue   # unchecked — leave the empty box
            ins = fitz.Rect(b.x0 + 0.7, b.y0 + 0.7, b.x1 - 0.7, b.y1 - 0.7)
            s = pg.new_shape(); s.draw_rect(ins); s.finish(fill=(1, 1, 1), color=None); s.commit(overlay=True)
            w, h = b.width, b.height
            p1 = fitz.Point(b.x0 + w * 0.22, b.y0 + h * 0.52)
            p2 = fitz.Point(b.x0 + w * 0.42, b.y0 + h * 0.74)
            p3 = fitz.Point(b.x0 + w * 0.80, b.y0 + h * 0.24)
            s2 = pg.new_shape(); s2.draw_line(p1, p2); s2.draw_line(p2, p3)
            s2.finish(color=(0, 0, 0), width=0.9, lineCap=1, lineJoin=1); s2.commit(overlay=True)
            total += 1
    doc.save(pdf_path, incremental=True, encryption=fitz.PDF_ENCRYPT_KEEP)
    doc.close()
    return total


# ─────────────────────────────────────────────────────────────────────────────
# SECTION 3 TABLE — auto-populated characteristics (added 2026-09-17)
# ─────────────────────────────────────────────────────────────────────────────
# Section 3 remains engineer-owned in the GOVERNED receive.py pipeline —
# CLAUDE.md's PRE-RUN REQUIREMENTS and Step 8 are unchanged: automation never
# writes a value to Section 3 on its own initiative, and receive.py does not
# call anything below. This exists for the OTHER real, recurring way Section
# 3 actually gets filled in practice: the engineer directly asking a Claude
# session with real device/file access to auto-populate the characteristic/
# tolerance table from a drawing's own numbered callouts. Every real instance
# of that in this codebase's history (PO26-1013/VendorH above all) did it by
# hand with ad hoc, unshared code — inconsistent column sizing, no enforced
# mm-stripping, no check that a "verbatim" drawing note actually was verbatim
# — exactly the "recurring pattern never folded into a shared helper" issue
# the 2026-09-11 gap analysis flagged as Priority 3 item #7. This is that
# fold-in, for Section 3 specifically, per QE's direct report.
#
# Cell map confirmed 2026-09-17 by opening the real blank template directly
# (templates/PRF template CONTROLS.xlsx) — not inferred, not carried over
# from the older/stale project-doc cell map, which had this one row off:
#   Row 6  D6:M6   — characteristic number labels ("1".."10", static,
#                     already present in the template — never written here)
#   Row 7  D7:M7   — "Criteria or Dim." (the dimension/description text)
#   Row 8  D8:M8   — Upper Limit
#   Row 9  D9:M9   — Lower Limit
#   Row 10 D10:M10 — Gaging Method
#   Rows 11-20     — per-unit measurements (engineer fills — never written)
#   Row 21         — array formulas (hard rule — never touched, any function)
# Column D is characteristic 1, column M is characteristic 10 — same
# left-to-right numbering convention already used by annotate_drawing_
# callouts() for the drawing's own bubbles (bubble N <-> column N), so a
# characteristic list built in drawing-bubble order lands in the right
# columns with no extra translation step.

_MM_BRACKET_RE = re.compile(r'\s*\[[^\[\]]*\]')


def strip_mm_bracket(text: str) -> str:
    """
    Remove a drawing dimension's bracketed millimeter cross-reference,
    e.g. '.984 [25.00]' -> '.984', 'M8X1.0-6H Thread Depth .118 [3.00]' ->
    'M8X1.0-6H Thread Depth .118'. Section 3 records inches only — the
    drawing's own mm reference is dropped, not converted or kept alongside.
    Only the bracketed group itself is removed; surrounding text (e.g. a
    trailing sheet reference like '(Sh3)') is left exactly as given. A
    string with no bracket, or an empty/None value, is returned unchanged.
    """
    if not text:
        return text
    return _MM_BRACKET_RE.sub('', text).strip()


def _sec3_text_lines(value) -> list:
    if value in (None, ''):
        return []
    return str(value).split('\n')


def _estimate_col_width(texts, min_width: float = 8.43, max_width: float = 42.0,
                         pad: float = 2.2) -> float:
    """
    Estimate the Excel column width (character units) needed to show the
    longest single line among `texts` without clipping. Character-count
    heuristic only — there is no LibreOffice/Excel renderer available in
    this environment to measure actual glyph widths, so this is an
    estimate, not a guarantee; spot-check the first real export, the same
    caveat this codebase attaches to every sizing/geometry guess made
    without a live renderer (e.g. check_row_text_clipping()'s own notice).
    Never returns less than `min_width` (the template's own default column
    width), so this only ever widens a column, never shrinks one below what
    the blank template already provides. Capped at `max_width` so one
    unusually long line can't blow out the whole table — anything longer
    than that wraps instead (see _estimate_row_height, which is computed
    against whatever width this function actually returns).
    """
    longest = 0
    for t in texts:
        for line in _sec3_text_lines(t):
            longest = max(longest, len(line))
    if longest == 0:
        return min_width
    return min(max(min_width, longest + pad), max_width)


def _estimate_row_height(texts, col_width: float, min_height: float = 14.7,
                          line_height: float = 15.0) -> float:
    """
    Estimate the row height needed for the tallest wrapped cell among
    `texts`, given the column width each will actually render in (same
    character-count heuristic and same "estimate, not exact" caveat as
    _estimate_col_width — see its docstring).
    """
    chars_per_line = max(1, int(col_width) - 1)
    tallest_lines = 1
    for t in texts:
        for line in _sec3_text_lines(t):
            n_wrapped = max(1, -(-len(line) // chars_per_line))  # ceil division
            tallest_lines = max(tallest_lines, n_wrapped)
    return max(min_height, tallest_lines * line_height)


def _normalize_ws(text: str) -> str:
    return re.sub(r'\s+', ' ', str(text or '')).strip().upper()


def fill_section3_table(ws, characteristics: list, autofit: bool = True) -> None:
    """
    Populate Section 3's characteristic/tolerance table (see the cell-map
    header above this function) from drawing-derived data, fixing three
    permanent formatting problems reported 2026-09-17 in one shared place
    instead of by hand on every PO:

      1. INCHES ONLY — every `criteria` string has its bracketed millimeter
         cross-reference stripped via strip_mm_bracket() before it's
         written. The caller supplies the criteria text exactly as read off
         the drawing (mm bracket and all); this function strips it, so it
         can never be forgotten on some characteristics and not others.

      2. MINIMUM AUTOFIT, CENTERED — when autofit=True (the default), each
         characteristic's own column is sized to the minimum width that
         fits its own longest line without clipping (see
         _estimate_col_width — a character-count estimate, spot-check the
         first real export), and rows 7 and 10 (Criteria / Gaging Method —
         the two free-text rows) are sized to the minimum height that fits
         whatever ended up needing to wrap at that width. Every cell this
         function writes is centered on both axes with wrap_text on,
         matching the 2026-09-17 fix applied to check_box()/uncheck_box()
         and to Section 6.2's B31 field — nothing here should drift to a
         corner of its own field. Pass autofit=False to only write values
         and leave existing column widths/row heights untouched.

      3. VERBATIM NOTES, ENFORCED DRAWING ORDER — for a fit-check /
         laser-etch / laser-mark characteristic, pass `verbatim_source`:
         the drawing's own note text for that callout, extracted
         character-for-character, never paraphrased or summarized. If
         `criteria` (after mm-stripping) does not appear as a substring of
         a whitespace-normalized, case-insensitive `verbatim_source`, this
         raises ValueError instead of silently writing a paraphrase — same
         "never guess, fail loud" posture as ChecklistCheckboxError
         elsewhere in this codebase. Drawing order is enforced structurally,
         not left to convention: `characteristics` must supply `number`
         1..N with no gaps or repeats, in that order — column D is always
         characteristic 1, matching the drawing's own bubble numbering (see
         annotate_drawing_callouts()) — or this raises ValueError before
         writing anything, so a caller can't accidentally hand characteristic
         3 to column D.

    `characteristics`: list of dicts, one per characteristic:
        number   (required) — 1-based position; the full list must form an
                 unbroken 1..N sequence (this is what makes drawing order a
                 structural guarantee rather than a convention)
        criteria (required) — dimension/description text; mm bracket okay,
                 stripped automatically
        upper, lower (optional) — numeric limits, OR the literal string
                 "High"/"Low" for a fit-check / GO-NO-GO / attribute
                 characteristic (per the real 2026-09-15 VendorH precedent:
                 Upper=High/Lower=Low for an attribute characteristic, kept
                 numeric for a dimensional one even when a fit-check rides
                 along with it — e.g. a thread depth with a mating-gauge
                 check stays numeric). NOTE, not fixed here: a text Upper/
                 Lower sitting in a row with numeric siblings makes the
                 row 21 In-Spec array formula misbehave, since Excel sorts
                 text above numbers — this is a pre-existing template
                 limitation and row 21 is never touched by any function in
                 this codebase; flag it for the engineer as always, don't
                 attempt to work around it here.
        gaging   (optional) — gaging method text
        verbatim_source (optional) — see point 3 above; omit for a plain
                 dimensional characteristic with no drawing note to match

    Never writes D6:M6 (characteristic number labels — already present in
    the template), D11:M20 (measurement rows — engineer fills), or D21:M21
    (array formulas — hard rule, no exceptions).
    """
    if not characteristics:
        return

    ordered = sorted(characteristics, key=lambda c: c.get("number", -1))
    expected_numbers = list(range(1, len(ordered) + 1))
    actual_numbers = [c.get("number") for c in ordered]
    if actual_numbers != expected_numbers:
        raise ValueError(
            f"fill_section3_table: characteristics must supply 'number' as an "
            f"unbroken 1..N sequence matching the drawing's own bubble order "
            f"(column D = characteristic 1) — got {actual_numbers}, expected "
            f"{expected_numbers}. Fix the caller's characteristic list rather "
            f"than reordering here; this function will not guess which "
            f"column a characteristic belongs in."
        )
    if len(ordered) > 10:
        raise ValueError(
            f"fill_section3_table: {len(ordered)} characteristics given, but "
            f"the template only has 10 columns (D..M) — split across multiple "
            f"PRFs/lots or confirm this is really one part before proceeding."
        )

    cleaned = []
    for c in ordered:
        criteria = strip_mm_bracket(c.get("criteria", ""))
        verbatim = c.get("verbatim_source")
        if verbatim:
            if _normalize_ws(criteria) not in _normalize_ws(verbatim):
                raise ValueError(
                    f"fill_section3_table: characteristic {c['number']}'s criteria "
                    f"text {criteria!r} is not a verbatim match against its own "
                    f"verbatim_source {verbatim!r} (whitespace/case-normalized "
                    f"comparison). Fit-check/laser-etch/laser-mark notes must be "
                    f"copied character-for-character off the drawing, not "
                    f"paraphrased or summarized — re-copy the note text exactly, "
                    f"or drop verbatim_source if this characteristic genuinely "
                    f"has no drawing note to match against."
                )
        cleaned.append({**c, "criteria": criteria})

    center_wrap = Alignment(horizontal='center', vertical='center', wrap_text=True)
    row_texts = {7: [], 8: [], 9: [], 10: []}

    for c in cleaned:
        col = get_column_letter(3 + c["number"])  # D=col 4=3+1, ... M=col13=3+10
        col_texts = [c.get("criteria", ""), c.get("gaging", ""),
                     str(c.get("upper", "")), str(c.get("lower", ""))]

        def _write(row, value):
            cell = ws[f"{col}{row}"]
            cell.value = value
            cell.font = Font(name='Calibri', size=11)
            cell.alignment = center_wrap

        _write(7, c.get("criteria", ""))
        if "upper" in c and c["upper"] not in (None, ""):
            _write(8, c["upper"])
        if "lower" in c and c["lower"] not in (None, ""):
            _write(9, c["lower"])
        if "gaging" in c and c["gaging"] not in (None, ""):
            _write(10, c["gaging"])

        row_texts[7].append(c.get("criteria", ""))
        row_texts[8].append(str(c.get("upper", "")))
        row_texts[9].append(str(c.get("lower", "")))
        row_texts[10].append(c.get("gaging", ""))

        if autofit:
            col_width = _estimate_col_width(col_texts)
            ws.column_dimensions[col].width = max(
                ws.column_dimensions[col].width or 0, col_width
            )

    if autofit:
        for row in (7, 10):  # the two free-text rows — 8/9 are short numeric/High-Low
            if not any(row_texts[row]):
                continue
            # Height depends on the width EACH cell actually renders at, not
            # a single shared width for the whole row — a wide column (e.g.
            # a long gaging-method description) fitting on one line must not
            # inflate the row height just because some other, narrower
            # column's own text happens to wrap. Take the max wrap-height
            # needed across each characteristic's own column.
            needed_height = 14.7
            for c in cleaned:
                col = get_column_letter(3 + c["number"])
                col_width = ws.column_dimensions[col].width or 8.43
                cell_text = ws[f"{col}{row}"].value
                if cell_text in (None, ""):
                    continue
                needed_height = max(needed_height, _estimate_row_height([cell_text], col_width))
            current = ws.row_dimensions[row].height or 14.7
            ws.row_dimensions[row].height = max(current, needed_height)
